# api/excel_import/processor.py
"""پردازش فایل اکسل"""

from django.utils import timezone
from openpyxl import load_workbook

from .staff_importer import import_staff
from .ward_importer import import_wards
from .holiday_importer import import_holidays
from .shiftcode_importer import import_shift_codes


def process_excel_upload(excel_upload) -> dict:
    """
    پردازش کامل فایل اکسل.
    """
    result = {
        'staff': None,
        'wards': None,
        'holidays': None,
        'shift_codes': None,
        'summary': {
            'total_created': 0,
            'total_updated': 0,
            'total_errors': 0,
        },
    }
    
    try:
        # بارگذاری فایل
        file_path = excel_upload.file.path
        wb = load_workbook(file_path, data_only=True)
        
        # ============================================================
        # ترتیب مهم: بخش‌ها اول، پرسنل آخر (وابسته به بخش)
        # ============================================================
        
        # ۱. کدهای شیفت
        if 'کدهای_شیفت' in wb.sheetnames:
            ws = wb['کدهای_شیفت']
            result['shift_codes'] = import_shift_codes(ws)
        
        # ۲. بخش‌ها
        if 'بخش‌ها' in wb.sheetnames:
            ws = wb['بخش‌ها']
            result['wards'] = import_wards(ws)
        
        # ۳. تعطیلات
        if 'تعطیلات' in wb.sheetnames:
            ws = wb['تعطیلات']
            result['holidays'] = import_holidays(ws)
        
        # ۴. پرسنل
        if 'پرسنل' in wb.sheetnames:
            ws = wb['پرسنل']
            result['staff'] = import_staff(ws)
        
        wb.close()
        
        # ============================================================
        # محاسبه خلاصه
        # ============================================================
        for section in ['staff', 'wards', 'holidays', 'shift_codes']:
            if result[section]:
                result['summary']['total_created'] += result[section].get('created', 0)
                result['summary']['total_updated'] += result[section].get('updated', 0)
                result['summary']['total_errors'] += len(result[section].get('errors', []))
        
        # ============================================================
        # تعیین وضعیت
        # ============================================================
        if result['summary']['total_errors'] == 0:
            excel_upload.status = 'success'
        elif result['summary']['total_created'] + result['summary']['total_updated'] > 0:
            excel_upload.status = 'partial'
        else:
            excel_upload.status = 'failed'
        
    except Exception as e:
        excel_upload.status = 'failed'
        excel_upload.error_message = str(e)
        result['error'] = str(e)
    
    # ذخیره نتیجه
    excel_upload.result_data = result
    excel_upload.processed_at = timezone.now()
    excel_upload.save()
    
    return result


def process_upload(excel_upload):
    """Wrapper برای فراخوانی از admin"""
    try:
        process_excel_upload(excel_upload)
    except Exception as e:
        print(f'❌ خطا در پردازش: {e}')