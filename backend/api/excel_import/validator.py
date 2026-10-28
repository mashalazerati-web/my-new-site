# api/excel_import/validator.py
"""اعتبارسنجی فایل اکسل"""

import os
from openpyxl import load_workbook


ALLOWED_EXTENSIONS = ['.xlsx']
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB


def validate_file(file_obj) -> tuple:
    """
    اعتبارسنجی فایل.
    Returns: (is_valid, error_message)
    """
    # فرمت
    if not file_obj.name.lower().endswith('.xlsx'):
        return False, 'فقط فایل با فرمت .xlsx مجاز است'
    
    # حجم
    if file_obj.size > MAX_FILE_SIZE:
        return False, f'حجم فایل نباید بیشتر از 5 مگابایت باشد (فعلی: {file_obj.size / 1024 / 1024:.1f}MB)'
    
    return True, None


def validate_workbook(file_path: str) -> tuple:
    """
    اعتبارسنجی ساختار فایل.
    Returns: (is_valid, error_message, sheets_found)
    """
    try:
        wb = load_workbook(file_path, read_only=True)
    except Exception as e:
        return False, f'فایل قابل خواندن نیست: {str(e)}', []
    
    # شیت‌های موردنیاز
    required_sheets = ['پرسنل', 'بخش‌ها', 'تعطیلات', 'کدهای_شیفت']
    found_sheets = wb.sheetnames
    
    # حداقل یک شیت باید موجود باشد
    available = [s for s in required_sheets if s in found_sheets]
    if not available:
        return False, (
            f'هیچ شیت معتبری یافت نشد. '
            f'شیت‌های موردانتظار: {", ".join(required_sheets)}'
        ), found_sheets
    
    wb.close()
    return True, None, available