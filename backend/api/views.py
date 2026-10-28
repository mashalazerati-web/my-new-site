# api/views.py
from rest_framework import viewsets, status, generics, serializers
from .models import (
    Ward, Profile, ShiftRequest, FinalSchedule, DraftSchedule,
    Announcement, Holiday, ShiftCode, SiteSettings, WardMonthlyNotes,
    MonthlyScheduleStatus
)
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.contrib.auth.models import User
from django.db import models
from .serializers import (
    UserSerializer, WardSerializer, ProfileSerializer,
    ShiftRequestSerializer, FinalScheduleSerializer, AnnouncementSerializer,
    ShiftCodeSerializer, SiteSettingsSerializer, DraftScheduleSerializer
)
import jdatetime
from .workflow import WorkflowEngine, WorkflowError
from .models import ScheduleStatus
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated, AllowAny
from django.http import HttpResponse
from django.http import FileResponse
from django.db.models import Q
import os
from django.conf import settings


# ========================================================
# ثابت‌ها
# ========================================================
ALL_SCHEDULE_TYPES = ['main', 'oncall', 'crisis', 'assistant']

SCHEDULE_TYPE_LABELS = {
    'main': 'اصلی',
    'oncall': 'آنکالی',
    'crisis': 'بحران',
    'assistant': 'کمک پرستار',
}

SCHEDULE_TYPE_LABELS_EMOJI = {
    'main': '🤖 اصلی',
    'oncall': '📞 آنکالی',
    'crisis': '🚨 بحران',
    'assistant': '🧑‍⚕️ کمک پرستار',
}


# ========================================================
# توابع کمکی برای تاریخ شمسی
# ========================================================

def get_persian_days_in_month(year, month):
    days_in_month = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    return days_in_month[month - 1]


def get_persian_weekday(year, month, day):
    """
    خروجی: 0=شنبه, 1=یکشنبه, ... 6=جمعه
    """
    try:
        jalali_date = jdatetime.date(year, month, day)
        return jalali_date.weekday()
    except Exception:
        return 0


# ========================================================
# تابع کمکی: مشتق کردن بولی‌ها از status
# ========================================================
def derive_booleans_from_status(server_status):
    """
    از روی status، بولی‌های نمایشی را برمی‌گرداند.
    این تنها منبع مشتق‌سازی بولی‌هاست.
    """
    return {
        'any_approved_by_head': server_status in (
            'PENDING_SUPERVISOR', 'PENDING_MATRON', 'LOCKED'
        ),
        'any_approved_by_supervisor': server_status in (
            'PENDING_MATRON', 'LOCKED'
        ),
        'is_locked': (server_status == 'LOCKED'),
        'is_returned_to_head': (server_status == 'RETURNED_TO_HEAD'),
        'is_returned_to_supervisor': (server_status == 'RETURNED_TO_SUPERVISOR'),
        'is_reset': (server_status == 'DRAFT'),
    }


def get_feedback_author(feedback_text, ward, role):
    """
    برگرداندن نام فردی که بازگردانی را نوشته است.
    """
    if not feedback_text:
        return None
    try:
        profile = Profile.objects.filter(ward=ward, role=role).first()
        if profile:
            return f"{profile.user.first_name} {profile.user.last_name}"
    except Exception:
        pass
    return None


def _get_profiles_for_schedule_type(ward, schedule_type, year=None, month=None):
    """
    ✅ تابع کمکی: گرفتن پرسنل بر اساس نوع برنامه.
    - assistant → فقط staff_category='assistant'
    - بقیه → فقط staff_category='nurse'
    - اعمال حذف‌های موقت
    """
    from .models import ScheduleStaffOverride

    if schedule_type == 'assistant':
        category_filter = Q(staff_category='assistant')
    else:
        category_filter = Q(staff_category='nurse')

    profiles = Profile.objects.filter(
        (Q(ward=ward) | Q(additional_wards=ward)) & category_filter
    ).exclude(
        role__in=['head', 'admin', 'supervisor']
    ).select_related('user').distinct()

    # ✅ اعمال حذف‌های موقت
    if year and month:
        excluded_staff_ids = set(
            ScheduleStaffOverride.objects.filter(
                ward=ward,
                year=int(year),
                month=int(month),
                schedule_type=schedule_type,
            ).values_list('staff_id', flat=True)
        )

        if excluded_staff_ids:
            profiles = profiles.exclude(user_id__in=excluded_staff_ids)

    return profiles


def _get_targets_for_schedule_type(ward, schedule_type):
    """
    ✅ تابع کمکی: گرفتن targetها بر اساس نوع برنامه.
    خروجی: (target_m, target_e, target_n, holiday_m, holiday_e, holiday_n)
    """
    if schedule_type in ('oncall', 'crisis', 'assistant'):
        from ai_module.models import SpecialScheduleConfig
        config = SpecialScheduleConfig.objects.filter(
            ward=ward, schedule_type=schedule_type
        ).first()

        if config:
            return (
                config.min_staff_morning,
                config.min_staff_evening,
                config.min_staff_night,
                config.holiday_min_staff_morning,
                config.holiday_min_staff_evening,
                config.holiday_min_staff_night,
            )
        else:
            if schedule_type == 'assistant':
                # پیش‌فرض: ۱ نفر هر شیفت
                return (1, 1, 1, 1, 1, 1)
            else:
                return (
                    ward.min_staff_morning,
                    ward.min_staff_evening,
                    ward.min_staff_night,
                    ward.holiday_min_staff_morning,
                    ward.holiday_min_staff_evening,
                    ward.holiday_min_staff_night,
                )
    else:
        # main
        return (
            ward.min_staff_morning,
            ward.min_staff_evening,
            ward.min_staff_night,
            ward.holiday_min_staff_morning,
            ward.holiday_min_staff_evening,
            ward.holiday_min_staff_night,
        )


# ========================================================
# تابع کمکی مشترک برای ساخت پاسخ برنامه
# ========================================================
def build_schedule_response(ward_id, year, month, source='final', mode='head', schedule_type='main'):
    """
    ساخت پاسخ برنامه ماهانه برای همه نقش‌ها.
    - برای main: ردیف سرپرستار اضافه می‌شود
    - برای oncall/crisis/assistant: ردیف سرپرستار اضافه نمی‌شود
    """
    from django.db.models import Q

    # محافظ
    try:
        year = int(year)
    except (ValueError, TypeError):
        year = jdatetime.date.today().year

    try:
        month = int(month)
    except (ValueError, TypeError):
        month = jdatetime.date.today().month + 1
        if month > 12:
            month = 1
            year += 1

    # 1. ward
    try:
        ward = Ward.objects.get(id=ward_id)
    except Ward.DoesNotExist:
        return Response({'error': 'بخش یافت نشد'}, status=404)

    # 2. site_settings
    site_settings = SiteSettings.objects.filter(is_active=True).first()

    # 3. profiles — ✅ با تابع کمکی
    profiles = list(_get_profiles_for_schedule_type(ward, schedule_type, year, month))

    profiles = sorted(
        profiles,
        key=lambda p: (
            int(p.position) if p.position and p.position.isdigit() else 999,
            -p.work_experience
        )
    )

    # 4. محاسبه روزها
    days_in_month = get_persian_days_in_month(year, month)
    first_day_index = get_persian_weekday(year, month, 1)
    weekdays = ['شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه']

    month_prefix = f"{year}-{str(month).zfill(2)}"

    # 5. FinalSchedule
    all_schedules = FinalSchedule.objects.filter(
        ward=ward,
        schedule_date__startswith=month_prefix,
        schedule_type=schedule_type
    ).select_related('assigned_staff', 'shift_code')

    # requests
    requests_dict = {}
    if source == 'staff_request' and schedule_type == 'main':
        shift_requests = ShiftRequest.objects.filter(
            ward=ward, request_date__startswith=month_prefix
        ).select_related('user', 'shift_code')
        for req in shift_requests:
            try:
                day = int(req.request_date.split('-')[2])
            except (ValueError, IndexError):
                continue
            if req.user.id not in requests_dict:
                requests_dict[req.user.id] = {}
            requests_dict[req.user.id][day] = req.shift_code.code if req.shift_code else '-'

    # head_draft
    draft_schedule_dict = {}
    if source == 'head_draft':
        draft_schedules = DraftSchedule.objects.filter(
            ward=ward,
            schedule_date__startswith=month_prefix,
            source='manual',
            schedule_type=schedule_type
        ).select_related('assigned_staff', 'shift_code')
        for ds in draft_schedules:
            try:
                day = int(ds.schedule_date.split('-')[2])
            except (ValueError, IndexError):
                continue
            if ds.assigned_staff.id not in draft_schedule_dict:
                draft_schedule_dict[ds.assigned_staff.id] = {}
            draft_schedule_dict[ds.assigned_staff.id][day] = ds

    # AI
    ai_schedule_dict = {}
    if source == 'ai_draft':
        from ai_module.models import AISchedule
        ai_schedule = AISchedule.objects.filter(
            ward=ward, year=year, month=month,
            schedule_type=schedule_type
        ).first()
        if ai_schedule and ai_schedule.ai_generated_data:
            for staff_item in ai_schedule.ai_generated_data.get('schedule_data', []):
                staff_id = staff_item.get('staff_id')
                if staff_id not in ai_schedule_dict:
                    ai_schedule_dict[staff_id] = {}
                for day_item in staff_item.get('days', []):
                    day = day_item.get('day')
                    shift_code = day_item.get('shift')
                    if shift_code:
                        try:
                            shift_obj = ShiftCode.objects.get(code=shift_code)
                            ai_schedule_dict[staff_id][day] = shift_obj
                        except ShiftCode.DoesNotExist:
                            continue

    # 6. تعطیلات
    month_holidays = Holiday.objects.filter(date__startswith=month_prefix).values('date', 'name')
    holidays_dict = {h['date']: h['name'] for h in month_holidays}

    # ============================================================
    # ✅ 6.5. پیدا کردن سرپرستار — فقط برای main
    # ============================================================
    head_profile = None
    head_user_id = None
    # ✅ فقط main سرپرستار داره
    if ward.parent is None and schedule_type == 'main':
        head_profile = Profile.objects.filter(ward=ward, role='head').first()
        if head_profile:
            head_user_id = head_profile.user.id

    table_data = []

    # ✅ ردیف سرپرستار (اول) — فقط برای main
    if head_profile:
        row_days = []
        total_m = total_e = total_n = total_h = 0
        total_off = 0

        for day in range(1, days_in_month + 1):
            day_index = get_persian_weekday(year, month, day)
            is_weekend = (day_index == 6)
            weekday_name = weekdays[day_index]

            date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
            is_holiday = date_key in holidays_dict
            holiday_name = holidays_dict.get(date_key, '')

            # ✅ منطق سرپرستار
            if is_holiday or is_weekend:
                shift = '-'
            else:
                shift = 'M'

            if head_user_id in requests_dict:
                requested = requests_dict[head_user_id].get(day)
                if requested in ['-', 'H', 'S', 'off']:
                    shift = '-'

            assigned_shifts = [shift] if shift != '-' else []

            if 'M' in shift: total_m += 1
            if 'E' in shift: total_e += 1
            if 'N' in shift: total_n += 1
            if 'H' in shift: total_h += 1
            if shift == '-': total_off += 1

            row_days.append({
                'day': day,
                'weekday': weekday_name,
                'is_weekend': is_weekend,
                'assigned_shifts': assigned_shifts,
                'approval_status': {
                    'approved_by_head': False,
                    'approved_by_supervisor': False,
                    'approved_by_matron': False,
                    'is_locked': False,
                    'returned_to_head': False,
                    'matron_feedback': None,
                    'supervisor_feedback': None,
                },
                'is_holiday': is_holiday,
                'holiday_name': holiday_name,
                'requested_shift': None,
            })

        table_data.append({
            'staff_id': head_user_id,
            'staff_name': f"{head_profile.user.first_name} {head_profile.user.last_name}".strip() or head_profile.user.username,
            'position_label': 'سرپرستار',
            'position_priority': head_profile.position,
            'work_experience': head_profile.work_experience,
            'employment_status': head_profile.employment_status or '',
            'shift_type_label': 'ثابت صبح',
            'days': row_days,
            'total_m': total_m,
            'total_e': total_e,
            'total_n': total_n,
            'total_h': total_h,
            'total_off': total_off,
            'is_head': True,
        })

    # ============================================================
    # 7. بقیه پرسنل
    # ============================================================
    for profile in profiles:
        # ✅ اگر این profile سرپرستار هست، از قبل اضافه شده
        if head_user_id and profile.user.id == head_user_id:
            continue

        row_days = []
        total_m = total_e = total_n = total_h = total_off = 0

        for day in range(1, days_in_month + 1):
            day_index = get_persian_weekday(year, month, day)
            is_weekend = (day_index == 6)
            weekday_name = weekdays[day_index]

            assigned_shifts = []
            approval_status = {
                'approved_by_head': False, 'approved_by_supervisor': False,
                'approved_by_matron': False, 'is_locked': False,
                'returned_to_head': False,
                'matron_feedback': None, 'supervisor_feedback': None,
                'matron_feedback_by': None, 'supervisor_feedback_by': None
            }

            # انتخاب منبع داده
            if source == 'staff_request':
                requested_shift = requests_dict.get(profile.user.id, {}).get(day)
                if requested_shift:
                    assigned_shifts = [requested_shift]
            elif source == 'head_draft':
                draft = draft_schedule_dict.get(profile.user.id, {}).get(day)
                if draft and draft.shift_code:
                    assigned_shifts = [draft.shift_code.code]
            elif source == 'ai_draft':
                shift_obj = ai_schedule_dict.get(profile.user.id, {}).get(day)
                if shift_obj:
                    assigned_shifts = [shift_obj.code]
            else:  # 'final'
                for schedule in all_schedules:
                    if (schedule.assigned_staff.id == profile.user.id
                            and int(schedule.schedule_date.split('-')[2]) == day):
                        if schedule.shift_code:
                            assigned_shifts.append(schedule.shift_code.code)
                            approval_status = {
                                'approved_by_head': schedule.approved_by_head,
                                'approved_by_supervisor': schedule.approved_by_supervisor,
                                'approved_by_matron': schedule.approved_by_matron,
                                'is_locked': schedule.is_locked,
                                'returned_to_head': schedule.returned_to_head,
                                'matron_feedback': schedule.matron_feedback,
                                'supervisor_feedback': schedule.supervisor_feedback,
                                'matron_feedback_by': get_feedback_author(schedule.matron_feedback, ward, 'admin'),
                                'supervisor_feedback_by': get_feedback_author(schedule.supervisor_feedback, ward, 'supervisor')
                            }
                            break

            for shift in assigned_shifts:
                if 'M' in shift: total_m += 1
                if 'E' in shift: total_e += 1
                if 'N' in shift: total_n += 1
                if 'H' in shift: total_h += 1
                if '-' in shift: total_off += 1

            row_days.append({
                'day': day,
                'weekday': weekday_name,
                'is_weekend': is_weekend,
                'assigned_shifts': assigned_shifts,
                'approval_status': approval_status,
                'is_holiday': False,
                'holiday_name': '',
                'requested_shift': None
            })

        shift_display_map = {
            'سه_نوبتی': 'سه نوبتی', 'ثابت_صبح': 'ثابت صبح', 'ثابت_عصر': 'ثابت عصر',
            'ثابت_شب': 'ثابت شب', 'دو_نوبتی_عصر_شب': 'دو نوبتی (عصر-شب)',
            'دو_نوبتی_صبح_شب': 'دو نوبتی (صبح-شب)', 'دو_نوبتی_صبح_عصر': 'دو نوبتی (صبح-عصر)',
        }
        raw_shift_type = profile.shift_type if hasattr(profile, 'shift_type') else '-'
        shift_type_label = shift_display_map.get(raw_shift_type, raw_shift_type)

        table_data.append({
            'staff_id': profile.user.id,
            'staff_name': f"{profile.user.first_name} {profile.user.last_name}",
            'position_label': profile.get_position_display(),
            'position_priority': profile.position,
            'work_experience': profile.work_experience,
            'employment_status': profile.employment_status if hasattr(profile, 'employment_status') else '-',
            'shift_type_label': shift_type_label,
            'days': row_days,
            'total_m': total_m, 'total_e': total_e, 'total_n': total_n,
            'total_h': total_h, 'total_off': total_off,
            'is_head': False,
        })

    # 8. تعطیلات روی روزها اعمال کن
    for row in table_data:
        for day_cell in row['days']:
            date_key = f"{year}-{str(month).zfill(2)}-{str(day_cell['day']).zfill(2)}"
            if date_key in holidays_dict:
                day_cell['is_holiday'] = True
                day_cell['holiday_name'] = holidays_dict[date_key]

    # 9. وضعیت
    if all_schedules.exists():
        server_status = all_schedules.first().status
    else:
        server_status = 'DRAFT'

    bools = derive_booleans_from_status(server_status)
    any_approved_by_head = bools['any_approved_by_head']
    any_approved_by_supervisor = bools['any_approved_by_supervisor']
    is_locked = bools['is_locked']
    is_returned_to_head = bools['is_returned_to_head']
    is_reset = bools['is_reset']

    # 10. امضاها
    supervisor_profile = Profile.objects.filter(ward=ward, role='supervisor').first()
    matron_profile = Profile.objects.filter(role='admin').first()

    head_signature_url = None
    supervisor_signature_url = None
    matron_signature_url = None

    if is_reset or is_returned_to_head:
        head_signature_url = None
        supervisor_signature_url = None
        matron_signature_url = None
    else:
        if any_approved_by_head and head_profile and head_profile.signature:
            head_signature_url = head_profile.signature.url
        if any_approved_by_supervisor and supervisor_profile and supervisor_profile.signature:
            supervisor_signature_url = supervisor_profile.signature.url
        if is_locked and matron_profile and matron_profile.signature:
            matron_signature_url = matron_profile.signature.url

    # 11. used_shift_codes
    used_shift_codes = set()
    for row in table_data:
        for day_cell in row['days']:
            if day_cell['assigned_shifts']:
                used_shift_codes.add(day_cell['assigned_shifts'][0])
    used_shift_codes_list = sorted(list(used_shift_codes))

    # 12. ward_notes
    ward_notes = ''
    monthly_note = WardMonthlyNotes.objects.filter(
        ward=ward, year=int(year), month=int(month)
    ).first()
    if monthly_note and monthly_note.notes:
        ward_notes = monthly_note.notes

    # 13. اهداف — ✅ با تابع کمکی
    (
        target_m, target_e, target_n,
        holiday_target_m, holiday_target_e, holiday_target_n
    ) = _get_targets_for_schedule_type(ward, schedule_type)

    # 14. Response
    return Response({
        "mode": "detail",
        "status": server_status,
        "schedule_type": schedule_type,
        "ward_info": {
            "id": ward.id,
            "name": ward.name,
            "head_name": ward.get_head_name(),
            "university_name": site_settings.university_name if site_settings else "",
            "university_logo": site_settings.university_logo.url if site_settings and site_settings.university_logo else None,
            "hospital_name": site_settings.hospital_name if site_settings else "",
            "hospital_logo": site_settings.hospital_logo.url if site_settings and site_settings.hospital_logo else None,
            "sub_ward_name": ward.sub_ward_name,
        },
        "table_data": table_data,
        "daily_m": [],
        "daily_e": [],
        "daily_n": [],
        "target_m": target_m,
        "target_e": target_e,
        "target_n": target_n,
        "holiday_target_m": holiday_target_m,
        "holiday_target_e": holiday_target_e,
        "holiday_target_n": holiday_target_n,
        "days_in_month": days_in_month,
        "weekdays": weekdays,
        "offset": first_day_index,
        "year": year,
        "month": month,
        "any_approved_by_head": any_approved_by_head,
        "any_approved_by_supervisor": any_approved_by_supervisor,
        "head_signature": head_signature_url,
        "supervisor_signature": supervisor_signature_url,
        "matron_signature": matron_signature_url,
        "used_shift_codes": used_shift_codes_list,
        "matron_feedback": all_schedules.first().matron_feedback if all_schedules.exists() else None,
        "supervisor_feedback": all_schedules.first().supervisor_feedback if all_schedules.exists() else None,
        "matron_feedback_by": get_feedback_author(
            all_schedules.first().matron_feedback if all_schedules.exists() else None, ward, 'admin'
        ),
        "supervisor_feedback_by": get_feedback_author(
            all_schedules.first().supervisor_feedback if all_schedules.exists() else None, ward, 'supervisor'
        ),
        "ward_notes": ward_notes
    })


# ========================================================
# ویو دریافت کدهای شیفت پویا
# ========================================================
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_shift_codes(request):
    try:
        profile = request.user.profile
        ward_id_param = request.query_params.get('ward_id')

        if ward_id_param:
            try:
                requested_ward_id = int(ward_id_param)
                user_wards = []
                if profile.ward:
                    user_wards.append(profile.ward.id)
                user_wards.extend([w.id for w in profile.additional_wards.all()])

                if requested_ward_id in user_wards:
                    ward = Ward.objects.get(id=requested_ward_id)
                else:
                    ward = profile.ward
            except (ValueError, TypeError, Ward.DoesNotExist):
                ward = profile.ward
        else:
            ward = profile.ward

        if not ward:
            return Response([])

        active_codes = ward.active_shift_codes.filter(is_active=True)
        serializer = ShiftCodeSerializer(active_codes, many=True)
        return Response(serializer.data)

    except Profile.DoesNotExist:
        return Response([])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_my_wards(request):
    try:
        profile = request.user.profile
        wards_list = []

        if profile.ward:
            wards_list.append({
                'id': profile.ward.id,
                'name': profile.ward.name,
                'is_primary': True,
            })

        for w in profile.additional_wards.all():
            wards_list.append({
                'id': w.id,
                'name': w.name,
                'is_primary': False,
            })

        return Response(wards_list)

    except Profile.DoesNotExist:
        return Response([])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_my_previous_schedules_wards(request):
    try:
        profile = request.user.profile
        user_ward_ids = []
        if profile.ward:
            user_ward_ids.append(profile.ward.id)
        user_ward_ids.extend([w.id for w in profile.additional_wards.all()])

        if not user_ward_ids:
            return Response([])

        year_param = request.query_params.get('year')
        month_param = request.query_params.get('month')

        schedules_qs = FinalSchedule.objects.filter(
            ward_id__in=user_ward_ids,
            assigned_staff=request.user,
            is_locked=True
        ).select_related('ward')

        if year_param and month_param:
            try:
                year = int(year_param)
                month = int(month_param)
                month_prefix = f"{year}-{str(month).zfill(2)}"
                schedules_qs = schedules_qs.filter(schedule_date__startswith=month_prefix)
            except (ValueError, TypeError):
                pass

        grouped = {}
        for s in schedules_qs:
            parts = s.schedule_date.split('-')
            if len(parts) != 3:
                continue
            y, m = int(parts[0]), int(parts[1])
            key = (s.ward_id, y, m)

            if key not in grouped:
                head_profile = Profile.objects.filter(ward=s.ward, role='head').first()
                head_name = f"{head_profile.user.first_name} {head_profile.user.last_name}" if head_profile else 'تعیین نشده'

                grouped[key] = {
                    'ward_id': s.ward_id,
                    'ward_name': s.ward.name,
                    'year': y,
                    'month': m,
                    'head_name': head_name,
                    'count': 0,
                }
            grouped[key]['count'] += 1

        result = list(grouped.values())
        result.sort(key=lambda x: (x['year'], x['month'], x['ward_name']), reverse=True)

        return Response(result)

    except Profile.DoesNotExist:
        return Response([])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_monthly_calendar(request):
    try:
        today = jdatetime.date.today()
        year = today.year
        month = today.month + 1
        if month > 12:
            month = 1
            year += 1

        year_param = request.query_params.get('year')
        month_param = request.query_params.get('month')
        if year_param:
            try:
                year = int(year_param)
            except Exception:
                pass
        if month_param:
            try:
                month = int(month_param)
            except Exception:
                pass

        try:
            profile = request.user.profile
        except Profile.DoesNotExist:
            return Response(
                {'error': 'پروفایل کاربر یافت نشد.'},
                status=status.HTTP_404_NOT_FOUND
            )

        ward_id_param = request.query_params.get('ward_id')

        if ward_id_param:
            user_wards = []
            if profile.ward:
                user_wards.append(profile.ward.id)
            user_wards.extend([w.id for w in profile.additional_wards.all()])

            try:
                requested_ward_id = int(ward_id_param)
            except (ValueError, TypeError):
                requested_ward_id = None

            if requested_ward_id and requested_ward_id in user_wards:
                try:
                    ward = Ward.objects.get(id=requested_ward_id)
                except Ward.DoesNotExist:
                    ward = profile.ward
            else:
                ward = profile.ward
        else:
            ward = profile.ward

        if not ward:
            return Response(
                {'error': 'شما به هیچ بخشی متصل نیستید.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        days_in_month = get_persian_days_in_month(year, month)
        month_prefix = f"{year}-{str(month).zfill(2)}"

        all_requests = ShiftRequest.objects.filter(
            ward_id=ward.id,
            user=request.user,
            request_date__startswith=month_prefix
        ).select_related('shift_code')

        is_user_locked_final = FinalSchedule.objects.filter(
            ward_id=ward.id,
            assigned_staff=request.user,
            schedule_date__startswith=month_prefix
        ).exists()

        is_user_locked_draft = DraftSchedule.objects.filter(
            ward_id=ward.id,
            assigned_staff=request.user,
            schedule_date__startswith=month_prefix,
            source='manual'
        ).exists()

        is_user_locked = is_user_locked_final or is_user_locked_draft

        requests_dict = {}
        for req in all_requests:
            try:
                day = int(req.request_date.split('-')[2])
                requests_dict[day] = req
            except (ValueError, IndexError):
                continue

        month_holidays = Holiday.objects.filter(
            date__startswith=month_prefix
        ).values('date', 'name')
        holidays_dict = {h['date']: h['name'] for h in month_holidays}

        days_data = []
        weekdays = ['شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه']
        first_day_weekday = get_persian_weekday(year, month, 1)

        for _ in range(first_day_weekday):
            days_data.append({
                'day': 0,
                'weekday': '',
                'is_weekend': False,
                'is_holiday': False,
                'holiday_name': '',
                'requests': []
            })

        for day in range(1, days_in_month + 1):
            date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"

            request_list = []

            if day in requests_dict:
                req = requests_dict[day]
                request_list.append({
                    'id': req.id,
                    'shift_type': req.shift_code.code if req.shift_code else '',
                    'is_approved': req.is_approved,
                    'source': 'request',
                    'approved_by_head': False,
                    'approved_by_supervisor': False,
                    'approved_by_matron': False
                })

            weekday_index = get_persian_weekday(year, month, day)
            is_weekend = (weekday_index == 6)
            is_holiday = date_key in holidays_dict
            holiday_name = holidays_dict.get(date_key, '')

            days_data.append({
                'day': day,
                'weekday': weekdays[weekday_index],
                'is_weekend': is_weekend,
                'is_holiday': is_holiday,
                'holiday_name': holiday_name,
                'requests': request_list
            })

        return Response({
            'year': year,
            'month': month,
            'days': days_data,
            'days_in_month': days_in_month,
            'is_locked': is_user_locked,
            'is_locked_final': is_user_locked_final,
            'is_locked_draft': is_user_locked_draft,
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response(
            {'error': str(e)},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


# ========================================================
# ویوست‌ها
# ========================================================

class UserViewSet(viewsets.ModelViewSet):
    queryset = User.objects.all()
    serializer_class = UserSerializer
    permission_classes = [IsAuthenticated]

    @action(detail=False, methods=['get'])
    def me(self, request):
        serializer = self.get_serializer(request.user)
        return Response(serializer.data)


class WardViewSet(viewsets.ModelViewSet):
    queryset = Ward.objects.all()
    serializer_class = WardSerializer
    permission_classes = [IsAuthenticated]


class ProfileViewSet(viewsets.ModelViewSet):
    queryset = Profile.objects.all()
    serializer_class = ProfileSerializer
    permission_classes = [IsAuthenticated]

    def create(self, request, *args, **kwargs):
        try:
            ward_id = request.data.get('ward')
            schedule_date_str = request.data.get('schedule_date')
            shift_input = request.data.get('shift_type')
            assigned_staff = request.data.get('assigned_staff')

            if not ward_id:
                return Response({"error": "شناسه بخش (ward) ارسال نشده است."}, status=status.HTTP_400_BAD_REQUEST)
            if not schedule_date_str:
                return Response({"error": "تاریخ نامعتبر است."}, status=status.HTTP_400_BAD_REQUEST)

            try:
                ward = Ward.objects.get(id=ward_id)
                user = User.objects.get(id=assigned_staff)
            except Ward.DoesNotExist:
                return Response({"error": "بخش مورد نظر یافت نشد."}, status=status.HTTP_404_NOT_FOUND)
            except User.DoesNotExist:
                return Response({"error": "پرستار مورد نظر یافت نشد."}, status=status.HTTP_404_NOT_FOUND)

            try:
                shift_code_obj = ShiftCode.objects.get(code__iexact=shift_input.upper())
            except ShiftCode.DoesNotExist:
                return Response({"error": f"کد شیفت '{shift_input}' در سیستم تعریف نشده است."}, status=status.HTTP_400_BAD_REQUEST)

            existing_schedule = FinalSchedule.objects.filter(
                ward=ward,
                schedule_date=schedule_date_str,
                assigned_staff=user
            ).first()

            if existing_schedule:
                existing_schedule.shift_code = shift_code_obj
                existing_schedule.save()
                serializer = self.get_serializer(existing_schedule)
                return Response(serializer.data, status=status.HTTP_200_OK)
            else:
                new_schedule = FinalSchedule.objects.create(
                    ward=ward,
                    schedule_date=schedule_date_str,
                    shift_code=shift_code_obj,
                    assigned_staff=user
                )
                serializer = self.get_serializer(new_schedule)
                return Response(serializer.data, status=status.HTTP_201_CREATED)

        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)


class ShiftRequestViewSet(viewsets.ModelViewSet):
    queryset = ShiftRequest.objects.all()
    serializer_class = ShiftRequestSerializer
    permission_classes = [IsAuthenticated]

    def _is_user_locked(self, user, ward, request_date):
        try:
            parts = request_date.split('-')
            year = int(parts[0])
            month = int(parts[1])
        except (ValueError, IndexError, AttributeError):
            return False

        month_prefix = f"{year}-{str(month).zfill(2)}"

        if FinalSchedule.objects.filter(
            ward=ward,
            assigned_staff=user,
            schedule_date__startswith=month_prefix
        ).exists():
            return True

        if DraftSchedule.objects.filter(
            ward=ward,
            assigned_staff=user,
            schedule_date__startswith=month_prefix,
            source='manual'
        ).exists():
            return True

        return False

    def perform_create(self, serializer):
        user = self.request.user
        profile = user.profile

        persian_date_str = serializer.validated_data['request_date']
        shift_input = self.request.data.get('shift_code') or self.request.data.get('shift_type')
        ward_id_from_request = self.request.data.get('ward')

        if ward_id_from_request:
            try:
                ward_id_int = int(ward_id_from_request)
            except (ValueError, TypeError):
                raise serializers.ValidationError({
                    "ward": "شناسه بخش نامعتبر است."
                })

            user_wards = []
            if profile.ward:
                user_wards.append(profile.ward.id)
            user_wards.extend([w.id for w in profile.additional_wards.all()])

            if ward_id_int not in user_wards:
                raise serializers.ValidationError({
                    "ward": "شما به این بخش دسترسی ندارید."
                })

            try:
                ward = Ward.objects.get(id=ward_id_int)
            except Ward.DoesNotExist:
                raise serializers.ValidationError({
                    "ward": "بخش مورد نظر یافت نشد."
                })
        else:
            ward = profile.ward

        if not ward:
            raise serializers.ValidationError({
                "ward": "شما به هیچ بخشی متصل نیستید."
            })

        if self._is_user_locked(user, ward, persian_date_str):
            raise serializers.ValidationError({
                "error": "برنامه شما توسط سرپرستار تنظیم شده و قابل ویرایش نیست."
            })

        try:
            shift_code_obj = ShiftCode.objects.get(code__iexact=shift_input)
        except ShiftCode.DoesNotExist:
            raise serializers.ValidationError({
                "shift_type": f"کد شیفت '{shift_input}' در کل سیستم تعریف نشده است."
            })

        if not ward.active_shift_codes.filter(id=shift_code_obj.id).exists():
            raise serializers.ValidationError({
                "shift_type": f"کد شیفت '{shift_input}' در سیستم وجود دارد اما برای بخش '{ward.name}' فعال نشده است. لطفاً به سرپرستار اطلاع دهید."
            })

        existing_request = ShiftRequest.objects.filter(
            user=user,
            ward=ward,
            request_date=persian_date_str
        ).first()

        if existing_request:
            existing_request.shift_code = shift_code_obj
            existing_request.is_approved = False
            existing_request.save()
        else:
            serializer.save(
                user=user,
                ward=ward,
                request_date=persian_date_str,
                shift_code=shift_code_obj,
                is_approved=False
            )

    def partial_update(self, request, *args, **kwargs):
        instance = self.get_object()
        user = request.user
        action = request.data.get('action')

        if not action:
            return Response(
                {'error': 'پارامتر action الزامی است.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        if action == 'delete':
            if instance.user != user:
                return Response(
                    {'error': 'شما اجازه حذف این درخواست را ندارید.'},
                    status=status.HTTP_403_FORBIDDEN
                )

            if instance.is_approved:
                return Response(
                    {'error': 'این درخواست تایید شده و قابل حذف نیست.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            if self._is_user_locked(user, instance.ward, instance.request_date):
                return Response(
                    {'error': 'برنامه شما توسط سرپرستار تنظیم شده و قابل ویرایش نیست.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            instance.delete()
            return Response({
                'status': 'deleted',
                'message': 'درخواست با موفقیت حذف شد.'
            })

        if action == 'approve':
            if user.profile.role != 'head':
                return Response(
                    {'error': 'فقط سرپرستار می‌تواند تایید کند.'},
                    status=status.HTTP_403_FORBIDDEN
                )

            instance.is_approved = True
            instance.save()
            return Response({
                'status': 'approved',
                'message': 'درخواست توسط سرپرستار تایید شد.'
            })

        if action == 'reject':
            if user.profile.role != 'head':
                return Response(
                    {'error': 'فقط سرپرستار می‌تواند رد کند.'},
                    status=status.HTTP_403_FORBIDDEN
                )

            instance.delete()
            return Response({
                'status': 'rejected',
                'message': 'درخواست توسط سرپرستار رد شد.'
            })

        return Response(
            {'error': 'action باید یکی از مقادیر "approve"، "reject" یا "delete" باشد.'},
            status=status.HTTP_400_BAD_REQUEST
        )


# ========================================================
# ویو برای برنامه ماهانه سرپرستار (Head Schedule)
# ========================================================
class HeadScheduleViewSet(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]

    def list(self, request):
        year_param = request.query_params.get('year')
        month_param = request.query_params.get('month')

        if year_param and month_param:
            try:
                year = int(year_param)
                month = int(month_param)
            except (ValueError, TypeError):
                today = jdatetime.date.today()
                year = today.year
                month = today.month + 1
                if month > 12:
                    month = 1
                    year += 1
        else:
            today = jdatetime.date.today()
            year = today.year
            month = today.month + 1
            if month > 12:
                month = 1
                year += 1

        ward_id = request.query_params.get('ward_id')
        source = request.query_params.get('source', 'final')

        if ward_id:
            return self._get_ward_schedule(request, ward_id, year, month, source)

        return self._get_ward_list(request, year, month)

    def _get_ward_schedule(self, request, ward_id, year, month, source='final'):
        schedule_type = request.query_params.get('schedule_type', 'main')

        return build_schedule_response(
            ward_id, year, month,
            source=source, mode='head',
            schedule_type=schedule_type
        )

    def _get_ward_list(self, request, year, month):
        user_profile = request.user.profile
        main_ward = user_profile.ward

        if not main_ward:
            return Response({
                "mode": "list",
                "wards": [],
                "year": year,
                "month": month,
                "error": "شما به هیچ بخش اصلی متصل نیستید."
            })

        all_wards = [main_ward] + list(main_ward.get_sub_wards())
        month_prefix = f"{year}-{str(month).zfill(2)}"

        from ai_module.models import AISchedule

        ward_list = []
        for w in all_wards:
            # ✅ چک وجود برنامه‌ها
            schedules_final = FinalSchedule.objects.filter(
                ward=w,
                schedule_date__startswith=month_prefix
            )

            schedules_draft = DraftSchedule.objects.filter(
                ward=w,
                schedule_date__startswith=month_prefix,
                source='manual'
            )

            schedules_ai = AISchedule.objects.filter(
                ward=w, year=year, month=month
            )

            # ✅ چک درخواست‌های شیفت
            shift_requests = ShiftRequest.objects.filter(
                ward=w,
                request_date__startswith=month_prefix
            )
            has_requests = shift_requests.exists()

            # ✅ چک پرسنل در بخش اصلی
            staff_count_main = Profile.objects.filter(
                ward=w, role='staff'
            ).count()

            # ✅ چک پرسنل در بخش‌های اضافی
            staff_count_additional = Profile.objects.filter(
                additional_wards=w, role='staff'
            ).count()

            total_staff = staff_count_main + staff_count_additional
            has_staff = total_staff > 0

            has_schedules = (
                schedules_final.exists() or
                schedules_draft.exists() or
                schedules_ai.exists()
            )

            has_any = has_schedules or has_requests or has_staff

            if not has_any:
                continue

            # ============================================================
            # ✅ وضعیت هر چهار نوع برنامه
            # ============================================================
            schedule_statuses = {}
            for st in ALL_SCHEDULE_TYPES:   # ✅ شامل assistant
                status_obj = MonthlyScheduleStatus.objects.filter(
                    ward=w, year=year, month=month, schedule_type=st
                ).first()

                has_final_st = FinalSchedule.objects.filter(
                    ward=w,
                    schedule_date__startswith=month_prefix,
                    schedule_type=st
                ).exists()

                has_ai_st = AISchedule.objects.filter(
                    ward=w, year=year, month=month, schedule_type=st
                ).exists()

                has_draft_st = DraftSchedule.objects.filter(
                    ward=w,
                    schedule_date__startswith=month_prefix,
                    source='manual',
                    schedule_type=st
                ).exists()

                has_schedule = has_final_st or has_ai_st or has_draft_st

                schedule_statuses[st] = {
                    'status': status_obj.status if status_obj else ('DRAFT' if has_schedule else None),
                    'status_display': status_obj.get_status_display_persian() if status_obj else (
                        '⏳ پیش‌نویس' if has_schedule else None
                    ),
                    'has_schedule': has_schedule,
                    'has_final': has_final_st,
                    'has_ai': has_ai_st,
                    'has_draft': has_draft_st,
                }

            # ✅ وضعیت اصلی
            main_status = schedule_statuses['main']['status'] or 'DRAFT'

            status_text = {
                'DRAFT': '⏳ در انتظار تنظیم و تایید سرپرستار',
                'PENDING_HEAD': '⏳ در انتظار تایید سرپرستار',
                'PENDING_SUPERVISOR': '⏳ در انتظار تایید میانی سوپروایزر',
                'PENDING_MATRON': '⏳ در انتظار تایید نهایی مترون',
                'LOCKED': '🔒 قفل شده (تایید نهایی)',
                'RETURNED_TO_HEAD': '🔄 بازگردانده شده به سرپرستار',
                'RETURNED_TO_SUPERVISOR': '🔄 بازگردانده شده به سوپروایزر'
            }.get(main_status, '⏳ در گردش کار')

            bools = derive_booleans_from_status(main_status)

            ward_list.append({
                'id': w.id,
                'name': w.name,
                'staff_count': total_staff,
                'head_name': w.get_head_name(),
                'status': main_status,
                'status_text': status_text,
                'schedule_statuses': schedule_statuses,
                'is_main': (w.id == main_ward.id),
                'is_sub': (w.parent is not None) if hasattr(w, 'parent') else False,
                'is_locked': bools['is_locked'],
                'is_approved_by_matron': main_status == 'LOCKED',
                'is_approved_by_supervisor': bools['any_approved_by_supervisor'],
                'is_approved_by_head': bools['any_approved_by_head'],
                'is_returned_to_head': bools['is_returned_to_head'],
                'has_matron_feedback': schedules_final.filter(matron_feedback__isnull=False).exists(),
                'target_m': w.min_staff_morning,
                'target_e': w.min_staff_evening,
                'target_n': w.min_staff_night,
                'holiday_target_m': w.holiday_min_staff_morning,
                'holiday_target_e': w.holiday_min_staff_evening,
                'holiday_target_n': w.holiday_min_staff_night,
                'has_requests': has_requests,
                'request_count': shift_requests.count(),
                'has_staff': has_staff,
                'has_schedules': has_schedules,
            })

        return Response({
            "mode": "list",
            "wards": ward_list,
            "year": year,
            "month": month
        })


# ========================================================
# ویو برای برنامه نهایی (Final Schedule)
# ========================================================
class FinalScheduleViewSet(viewsets.ModelViewSet):
    queryset = FinalSchedule.objects.all()
    serializer_class = FinalScheduleSerializer
    permission_classes = [IsAuthenticated]

    def create(self, request, *args, **kwargs):
        try:
            ward_id = request.data.get('ward')
            schedule_date_str = request.data.get('schedule_date')
            shift_input = request.data.get('shift_type')
            assigned_staff = request.data.get('assigned_staff')

            if not ward_id:
                return Response({"error": "شناسه بخش (ward) ارسال نشده است."}, status=status.HTTP_400_BAD_REQUEST)
            if not schedule_date_str:
                return Response({"error": "تاریخ نامعتبر است."}, status=status.HTTP_400_BAD_REQUEST)

            try:
                ward = Ward.objects.get(id=ward_id)
                user = User.objects.get(id=assigned_staff)
            except Ward.DoesNotExist:
                return Response({"error": "بخش مورد نظر یافت نشد."}, status=status.HTTP_404_NOT_FOUND)
            except User.DoesNotExist:
                return Response({"error": "پرستار مورد نظر یافت نشد."}, status=status.HTTP_404_NOT_FOUND)

            try:
                shift_code_obj = ShiftCode.objects.get(code__iexact=shift_input.upper())
            except ShiftCode.DoesNotExist:
                return Response({"error": f"کد شیفت '{shift_input}' در سیستم تعریف نشده است."}, status=status.HTTP_400_BAD_REQUEST)

            existing_schedule = FinalSchedule.objects.filter(
                ward=ward,
                schedule_date=schedule_date_str,
                assigned_staff=user
            ).first()

            if existing_schedule:
                existing_schedule.shift_code = shift_code_obj
                existing_schedule.save()
                serializer = self.get_serializer(existing_schedule)
                return Response(serializer.data, status=status.HTTP_200_OK)
            else:
                new_schedule = FinalSchedule.objects.create(
                    ward=ward,
                    schedule_date=schedule_date_str,
                    shift_code=shift_code_obj,
                    assigned_staff=user
                )
                serializer = self.get_serializer(new_schedule)
                return Response(serializer.data, status=status.HTTP_201_CREATED)

        except Exception as e:
            print(f"❌ خطای بحرانی در ایجاد برنامه: {str(e)}")
            import traceback
            traceback.print_exc()
            return Response({"error": f"خطای داخلی سرور: {str(e)}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['get'])
    def head_schedule(self, request):
        return Response({'error': 'این متد منسوخ شده است. از HeadScheduleViewSet استفاده کنید.'}, status=400)

    # ============================================================
    # ✅ داشبورد سوپروایزر — چهار برنامه‌ای
    # ============================================================
    @action(detail=False, methods=['get'])
    def supervisor_dashboard(self, request):
        if request.user.profile.role != 'supervisor':
            return Response({'error': 'دسترسی غیرمجاز. این صفحه فقط مخصوص سوپروایزر است'}, status=403)

        year_param = request.query_params.get('year')
        month_param = request.query_params.get('month')

        if not year_param or not month_param:
            today = jdatetime.date.today()
            year = today.year
            month = today.month + 1
            if month > 12:
                month = 1
                year += 1
        else:
            try:
                year = int(year_param)
                month = int(month_param)
            except Exception:
                today = jdatetime.date.today()
                year = today.year
                month = today.month + 1
                if month > 12:
                    month = 1
                    year += 1

        month_prefix = f"{year}-{str(month).zfill(2)}"
        wards = request.user.supervised_wards.all()

        result = []

        for ward in wards:
            # ✅ شامل assistant
            for schedule_type in ALL_SCHEDULE_TYPES:
                all_schedules = FinalSchedule.objects.filter(
                    ward=ward,
                    schedule_date__startswith=month_prefix,
                    schedule_type=schedule_type
                ).select_related('assigned_staff', 'shift_code')

                if not all_schedules.exists():
                    continue

                total_m = 0
                total_e = 0
                total_n = 0
                staff_names = set()

                for schedule in all_schedules:
                    if schedule.assigned_staff:
                        staff_names.add(
                            f"{schedule.assigned_staff.first_name} {schedule.assigned_staff.last_name}"
                        )

                    if schedule.shift_code:
                        code = schedule.shift_code.code.upper()
                        if 'M' in code:
                            total_m += 1
                        if 'E' in code:
                            total_e += 1
                        if 'N' in code:
                            total_n += 1

                server_status = all_schedules.first().status

                status_text = {
                    'DRAFT': '⏳ در انتظار تنظیم و تایید سرپرستار',
                    'PENDING_HEAD': '⏳ در انتظار تایید سرپرستار',
                    'PENDING_SUPERVISOR': '⏳ در انتظار تایید میانی سوپروایزر',
                    'PENDING_MATRON': '⏳ در انتظار تایید نهایی مترون',
                    'LOCKED': '🔒 قفل شده (تایید نهایی)',
                    'RETURNED_TO_HEAD': '🔄 بازگردانده شده به سرپرستار',
                    'RETURNED_TO_SUPERVISOR': '🔄 بازگردانده شده به سوپروایزر'
                }.get(server_status, '⏳ در گردش کار')

                schedule_type_label = SCHEDULE_TYPE_LABELS_EMOJI.get(schedule_type, '🤖 اصلی')

                # ✅ استفاده از تابع کمکی برای target
                target_m, target_e, target_n, _, _, _ = _get_targets_for_schedule_type(
                    ward, schedule_type
                )

                result.append({
                    'id': ward.id,
                    'name': ward.name,
                    'head_name': ward.get_head_name(),
                    'month': month,
                    'year': year,
                    'schedule_type': schedule_type,
                    'schedule_type_label': schedule_type_label,
                    'staff_count': len(staff_names),
                    'staff_names': list(staff_names),
                    'total_m': total_m,
                    'total_e': total_e,
                    'total_n': total_n,
                    'target_m': target_m,
                    'target_e': target_e,
                    'target_n': target_n,
                    'schedule_ids': [s.id for s in all_schedules],
                    'status': server_status,
                    'status_text': status_text,
                })

        return Response(result)

    # ============================================================
    # ✅ داشبورد مترون — چهار برنامه‌ای
    # ============================================================
    @action(detail=False, methods=['get'])
    def matron_dashboard(self, request):
        if request.user.profile.role != 'admin':
            return Response({'error': 'دسترسی غیرمجاز. این صفحه فقط مخصوص مترون است.'}, status=403)

        year_param = request.query_params.get('year')
        month_param = request.query_params.get('month')

        if not year_param or not month_param:
            today = jdatetime.date.today()
            year = today.year
            month = today.month + 1
            if month > 12:
                month = 1
                year += 1
        else:
            try:
                year = int(year_param)
                month = int(month_param)
            except Exception:
                today = jdatetime.date.today()
                year = today.year
                month = today.month + 1
                if month > 12:
                    month = 1
                    year += 1

        month_prefix = f"{year}-{str(month).zfill(2)}"
        wards = Ward.objects.all()
        result = []

        for ward in wards:
            # ✅ شامل assistant
            for schedule_type in ALL_SCHEDULE_TYPES:
                schedules = FinalSchedule.objects.filter(
                    ward=ward,
                    schedule_date__startswith=month_prefix,
                    schedule_type=schedule_type
                ).select_related('assigned_staff', 'shift_code')

                if not schedules.exists():
                    continue

                staff_names = list(set([
                    s.assigned_staff.first_name + " " + s.assigned_staff.last_name
                    for s in schedules
                ]))

                supervisor_name = (
                    ward.supervisor.first_name + " " + ward.supervisor.last_name
                    if ward.supervisor else "تعیین نشده"
                )

                server_status = schedules.first().status

                status_text = {
                    'DRAFT': '⏳ در انتظار تنظیم و تایید سرپرستار',
                    'PENDING_HEAD': '⏳ در انتظار تایید سرپرستار',
                    'PENDING_SUPERVISOR': '⏳ در انتظار تایید میانی سوپروایزر',
                    'PENDING_MATRON': '⏳ در انتظار تایید نهایی مترون',
                    'LOCKED': '🔒 قفل شده (تایید نهایی)',
                    'RETURNED_TO_HEAD': '🔄 بازگردانده شده به سرپرستار',
                    'RETURNED_TO_SUPERVISOR': '🔄 بازگردانده شده به سوپروایزر'
                }.get(server_status, '⏳ در گردش کار')

                schedule_type_label = SCHEDULE_TYPE_LABELS_EMOJI.get(schedule_type, '🤖 اصلی')

                result.append({
                    'id': ward.id,
                    'name': ward.name,
                    'head_name': ward.get_head_name(),
                    'supervisor_name': supervisor_name,
                    'month': month,
                    'year': year,
                    'schedule_type': schedule_type,
                    'schedule_type_label': schedule_type_label,
                    'staff_count': len(staff_names),
                    'staff_names': staff_names,
                    'schedule_ids': [s.id for s in schedules],
                    'status': server_status,
                    'status_text': status_text,
                })

        return Response(result)

    # ============================================================
    # ✅ متدهای نمایش برنامه برای سوپروایزر و مترون
    # ============================================================
    @action(detail=False, methods=['get'])
    def supervisor_schedule(self, request):
        year = request.query_params.get('year')
        month = request.query_params.get('month')
        ward_id = request.query_params.get('ward_id')
        schedule_type = request.query_params.get('schedule_type', 'main')

        print(f"🔍 [supervisor_schedule] year={year}, month={month}, ward_id={ward_id}, schedule_type={schedule_type}")

        if not ward_id:
            return Response({'error': 'شناسه بخش (ward_id) لازم است'}, status=400)

        return build_schedule_response(
            ward_id, year, month,
            source='final', mode='supervisor',
            schedule_type=schedule_type
        )

    @action(detail=False, methods=['get'])
    def matron_schedule(self, request):
        year = request.query_params.get('year')
        month = request.query_params.get('month')
        ward_id = request.query_params.get('ward_id')
        schedule_type = request.query_params.get('schedule_type', 'main')

        if not ward_id:
            return Response({'error': 'شناسه بخش (ward_id) لازم است'}, status=400)

        return build_schedule_response(
            ward_id, year, month,
            source='final', mode='matron',
            schedule_type=schedule_type
        )


# ========================================================
# ویو ذخیره برنامه (Draft)
# ========================================================
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def save_schedule(request):
    try:
        ward_id = request.data.get('ward_id')
        schedule_data = request.data.get('data')
        year = request.data.get('year')
        month = request.data.get('month')
        ward_notes = request.data.get('ward_notes', '')
        schedule_type = request.data.get('schedule_type', 'main')

        if not ward_id or not schedule_data or not year or not month:
            return Response({'error': 'شناسه بخش، دیتای برنامه، سال و ماه الزامی است.'}, status=400)

        try:
            ward = Ward.objects.get(id=ward_id)
        except Ward.DoesNotExist:
            return Response({'error': 'بخش مورد نظر یافت نشد.'}, status=404)

        month_prefix = f"{year}-{str(month).zfill(2)}"

        # ✅ ۱. ذخیره ward_notes
        if ward_notes:
            WardMonthlyNotes.objects.update_or_create(
                ward=ward,
                year=int(year),
                month=int(month),
                defaults={
                    'notes': ward_notes,
                    'updated_by': request.user
                }
            )

        # ✅ ۲. پاک کردن DraftSchedule قبلی
        DraftSchedule.objects.filter(
            ward=ward,
            schedule_date__startswith=month_prefix,
            source='manual',
            schedule_type=schedule_type
        ).delete()

        # ✅ ۳. آماده‌سازی داده‌ها
        shift_codes_map = {
            sc.code: sc for sc in ShiftCode.objects.all()
        }

        staff_ids = [row['staff_id'] for row in schedule_data]
        users_map = {
            u.id: u for u in User.objects.filter(id__in=staff_ids)
        }

        # ✅ ۴. ساخت لیست برای bulk_create
        new_drafts = []

        for staff_row in schedule_data:
            staff_id = staff_row['staff_id']
            user = users_map.get(staff_id)

            if not user:
                continue

            for day_cell in staff_row['days']:
                day = day_cell['day']
                shift_code_str = day_cell['shift']

                shift_code_obj = shift_codes_map.get(shift_code_str)
                if not shift_code_obj:
                    continue

                schedule_date = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"

                new_drafts.append(
                    DraftSchedule(
                        ward=ward,
                        schedule_date=schedule_date,
                        shift_code=shift_code_obj,
                        assigned_staff=user,
                        source='manual',
                        schedule_type=schedule_type,
                        created_by=request.user
                    )
                )

        # ✅ ۵. درج گروهی
        if new_drafts:
            DraftSchedule.objects.bulk_create(
                new_drafts,
                batch_size=500,
                ignore_conflicts=True,
            )

        return Response({
            'status': 'success',
            'message': f'{len(new_drafts)} رکورد ذخیره شد.',
            'count': len(new_drafts),
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': f'خطای ذخیره‌سازی: {str(e)}'}, status=500)


# ========================================================
# ویو انتقال از DraftSchedule به FinalSchedule
# ========================================================
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def finalize_schedule(request):
    try:
        ward_id = request.data.get('ward_id')
        year = request.data.get('year')
        month = request.data.get('month')
        ward_notes = request.data.get('ward_notes', '')
        schedule_type = request.data.get('schedule_type', 'main')

        if not ward_id or not year or not month:
            return Response({'error': 'شناسه بخش، سال و ماه الزامی است.'}, status=400)

        try:
            ward = Ward.objects.get(id=ward_id)
        except Ward.DoesNotExist:
            return Response({'error': 'بخش مورد نظر یافت نشد.'}, status=404)

        month_prefix = f"{year}-{str(month).zfill(2)}"

        # ✅ ۱. پاک کردن FinalSchedule قبلی
        FinalSchedule.objects.filter(
            ward=ward,
            schedule_date__startswith=month_prefix,
            schedule_type=schedule_type
        ).delete()

        # ✅ ۲. خواندن همه Draft
        draft_schedules = list(DraftSchedule.objects.filter(
            ward=ward,
            schedule_date__startswith=month_prefix,
            source='manual',
            schedule_type=schedule_type
        ).select_related('assigned_staff', 'shift_code'))

        if not draft_schedules:
            return Response({'error': 'برنامه پیش‌نویس یافت نشد.'}, status=404)

        # ✅ ۳. ساخت لیست FinalSchedule
        new_schedules = []
        for ds in draft_schedules:
            new_schedules.append(
                FinalSchedule(
                    ward=ward,
                    schedule_date=ds.schedule_date,
                    shift_code=ds.shift_code,
                    assigned_staff=ds.assigned_staff,
                    schedule_type=schedule_type,
                    ward_notes=ward_notes,
                    status=ScheduleStatus.DRAFT,
                    approved_by_head=False,
                    approved_by_supervisor=False,
                    approved_by_matron=False,
                    is_locked=False,
                    returned_to_head=False,
                )
            )

        # ✅ ۴. درج گروهی
        FinalSchedule.objects.bulk_create(
            new_schedules,
            batch_size=500,
        )

        # ✅ ۵. ذخیره ward_notes
        if ward_notes:
            WardMonthlyNotes.objects.update_or_create(
                ward=ward,
                year=int(year),
                month=int(month),
                defaults={
                    'notes': ward_notes,
                    'updated_by': request.user
                }
            )

        return Response({
            'message': 'برنامه با موفقیت به FinalSchedule منتقل شد.',
            'count': len(new_schedules)
        }, status=200)

    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': f'خطای داخلی: {str(e)}'}, status=500)


# ========================================================
# ویو اخبار
# ========================================================
class AnnouncementListView(generics.ListAPIView):
    serializer_class = AnnouncementSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Announcement.objects.filter(
            is_active=True
        ).prefetch_related('attachments').order_by('-created_at')


# ========================================================
# کلاس TransitionStatusView - Endpoint واحد گردش کار
# ========================================================
class TransitionStatusView(APIView):
    """
    Endpoint واحد گردش کار — تغییر وضعیت برنامه.
    با بهینه‌سازی bulk_update و bulk_create.
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        try:
            ward_id = request.data.get('ward_id')
            year = request.data.get('year')
            month = request.data.get('month')
            action = request.data.get('action')
            reason = request.data.get('reason', '')
            schedule_type = request.data.get('schedule_type', 'main')

            # ============================================================
            # ۱. اعتبارسنجی
            # ============================================================
            if not ward_id or not year or not month or not action:
                return Response(
                    {'error': 'ward_id، year، month و action الزامی هستند.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # ✅ شامل assistant
            if schedule_type not in ALL_SCHEDULE_TYPES:
                return Response(
                    {'error': 'نوع برنامه نامعتبر است.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            try:
                ward = Ward.objects.get(id=ward_id)
            except Ward.DoesNotExist:
                return Response(
                    {'error': 'بخش مورد نظر یافت نشد.'},
                    status=status.HTTP_404_NOT_FOUND
                )

            actor_role = request.user.profile.role
            month_prefix = f"{year}-{str(month).zfill(2)}"

            # ============================================================
            # ۲. خواندن یا ساخت MonthlyScheduleStatus
            # ============================================================
            schedule_status_obj = MonthlyScheduleStatus.objects.filter(
                ward=ward,
                year=int(year),
                month=int(month),
                schedule_type=schedule_type
            ).first()

            # ✅ اگر نبود و action=approve با role=head → انتقال از Draft به Final
            if not schedule_status_obj and action == 'approve' and actor_role == 'head':
                draft_schedules = list(DraftSchedule.objects.filter(
                    ward=ward,
                    schedule_date__startswith=month_prefix,
                    source='manual',
                    schedule_type=schedule_type
                ).select_related('assigned_staff', 'shift_code'))

                if draft_schedules:
                    # پاک کردن Final قبلی
                    FinalSchedule.objects.filter(
                        ward=ward,
                        schedule_date__startswith=month_prefix,
                        schedule_type=schedule_type
                    ).delete()

                    # ✅ bulk_create برای Final
                    new_schedules = []
                    for ds in draft_schedules:
                        new_schedules.append(
                            FinalSchedule(
                                ward=ward,
                                schedule_date=ds.schedule_date,
                                shift_code=ds.shift_code,
                                assigned_staff=ds.assigned_staff,
                                schedule_type=schedule_type,
                                ward_notes=reason or '',
                                status=ScheduleStatus.DRAFT,
                                approved_by_head=False,
                                approved_by_supervisor=False,
                                approved_by_matron=False,
                                is_locked=False,
                                returned_to_head=False,
                            )
                        )

                    FinalSchedule.objects.bulk_create(
                        new_schedules,
                        batch_size=500,
                    )

                    # پاک کردن Draft
                    DraftSchedule.objects.filter(
                        ward=ward,
                        schedule_date__startswith=month_prefix,
                        source='manual',
                        schedule_type=schedule_type
                    ).delete()

                    # ساخت MonthlyScheduleStatus
                    schedule_status_obj = MonthlyScheduleStatus.objects.create(
                        ward=ward,
                        year=int(year),
                        month=int(month),
                        schedule_type=schedule_type,
                        status=ScheduleStatus.DRAFT,
                        updated_by=request.user
                    )
                else:
                    return Response(
                        {'error': 'برنامه‌ای برای این بخش، ماه و نوع برنامه یافت نشد.'},
                        status=status.HTTP_404_NOT_FOUND
                    )

            elif not schedule_status_obj:
                return Response(
                    {'error': f'برنامه {schedule_type} برای این بخش و ماه یافت نشد.'},
                    status=status.HTTP_404_NOT_FOUND
                )

            current_status = schedule_status_obj.status

            # ============================================================
            # ۳. تغییر وضعیت از طریق WorkflowEngine
            # ============================================================
            try:
                if action == 'approve':
                    new_status = WorkflowEngine.approve(
                        current_status, actor_role, schedule_type
                    )
                elif action == 'reject':
                    new_status, reason = WorkflowEngine.reject(
                        current_status, actor_role, reason, schedule_type
                    )
                elif action == 'reset':
                    new_status = WorkflowEngine.reset(
                        current_status, actor_role, schedule_type
                    )
                else:
                    raise WorkflowError("عملیات نامعتبر است.")

                # ✅ به‌روزرسانی MonthlyScheduleStatus
                schedule_status_obj.status = new_status
                schedule_status_obj.updated_by = request.user

                if action == 'reject':
                    if actor_role == 'supervisor':
                        schedule_status_obj.supervisor_feedback = reason or ''
                    elif actor_role == 'admin':
                        schedule_status_obj.matron_feedback = reason or ''
                elif action == 'reset':
                    schedule_status_obj.supervisor_feedback = ''
                    schedule_status_obj.matron_feedback = ''

                schedule_status_obj.save()

                # ============================================================
                # ۴. ✅ به‌روزرسانی FinalSchedule با bulk_update
                # ============================================================
                final_schedules = list(FinalSchedule.objects.filter(
                    ward=ward,
                    schedule_date__startswith=month_prefix,
                    schedule_type=schedule_type
                ))

                # ✅ محاسبه مقادیر جدید بر اساس status
                new_approved_by_head = new_status in (
                    ScheduleStatus.PENDING_SUPERVISOR,
                    ScheduleStatus.PENDING_MATRON,
                    ScheduleStatus.LOCKED,
                )
                new_approved_by_supervisor = new_status in (
                    ScheduleStatus.PENDING_MATRON,
                    ScheduleStatus.LOCKED,
                )
                new_approved_by_matron = new_status == ScheduleStatus.LOCKED
                new_is_locked = new_status == ScheduleStatus.LOCKED
                new_returned_to_head = new_status == ScheduleStatus.RETURNED_TO_HEAD

                # ✅ آماده‌سازی لیست برای bulk_update
                schedules_to_update = []
                for schedule in final_schedules:
                    schedule.status = new_status
                    schedule.approved_by_head = new_approved_by_head
                    schedule.approved_by_supervisor = new_approved_by_supervisor
                    schedule.approved_by_matron = new_approved_by_matron
                    schedule.is_locked = new_is_locked
                    schedule.returned_to_head = new_returned_to_head

                    if action == 'reject':
                        if actor_role == 'supervisor':
                            schedule.supervisor_feedback = reason or ''
                        elif actor_role == 'admin':
                            schedule.matron_feedback = reason or ''
                    elif action == 'reset':
                        schedule.supervisor_feedback = ''
                        schedule.matron_feedback = ''

                    schedules_to_update.append(schedule)

                # ✅ یک کوئری UPDATE به‌جای N کوئری
                if schedules_to_update:
                    FinalSchedule.objects.bulk_update(
                        schedules_to_update,
                        fields=[
                            'status',
                            'approved_by_head',
                            'approved_by_supervisor',
                            'approved_by_matron',
                            'is_locked',
                            'returned_to_head',
                            'supervisor_feedback',
                            'matron_feedback',
                        ],
                        batch_size=500,
                    )

                # ============================================================
                # ۵. ✅ اگر reset → ساخت Draft با bulk_create
                # ============================================================
                if new_status == ScheduleStatus.DRAFT:
                    # پاک کردن Draft قبلی
                    DraftSchedule.objects.filter(
                        ward=ward,
                        schedule_date__startswith=month_prefix,
                        source='manual',
                        schedule_type=schedule_type
                    ).delete()

                    # ✅ bulk_create برای Draft
                    new_drafts = []
                    for schedule in final_schedules:
                        new_drafts.append(
                            DraftSchedule(
                                ward=schedule.ward,
                                schedule_date=schedule.schedule_date,
                                shift_code=schedule.shift_code,
                                assigned_staff=schedule.assigned_staff,
                                source='manual',
                                schedule_type=schedule_type,
                                created_by=request.user
                            )
                        )

                    if new_drafts:
                        DraftSchedule.objects.bulk_create(
                            new_drafts,
                            batch_size=500,
                        )

                # ============================================================
                # ۶. ساخت پیام موفق
                # ============================================================
                type_label = SCHEDULE_TYPE_LABELS.get(schedule_type, schedule_type)

                return Response({
                    'status': 'success',
                    'new_status': new_status,
                    'schedule_type': schedule_type,
                    'message': f'وضعیت برنامه {type_label} با موفقیت به {new_status} تغییر کرد.'
                })

            except WorkflowError as e:
                return Response(
                    {'error': str(e)},
                    status=status.HTTP_400_BAD_REQUEST
                )

        except Exception as e:
            import traceback
            traceback.print_exc()
            return Response(
                {'error': f'خطای داخلی: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


# ========================================================
# کلاس MonthlyStatusView
# ========================================================
class MonthlyStatusView(APIView):
    """
    دریافت وضعیت هر چهار نوع برنامه برای یک بخش در یک ماه.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        ward_id = request.query_params.get('ward_id')
        year = request.query_params.get('year')
        month = request.query_params.get('month')

        if not ward_id or not year or not month:
            return Response(
                {'error': 'ward_id، year و month الزامی هستند.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            ward = Ward.objects.get(id=ward_id)
        except Ward.DoesNotExist:
            return Response(
                {'error': 'بخش مورد نظر یافت نشد.'},
                status=status.HTTP_404_NOT_FOUND
            )

        result = {}

        # ✅ شامل assistant
        for schedule_type in ALL_SCHEDULE_TYPES:
            # دریافت وضعیت
            status_obj = MonthlyScheduleStatus.objects.filter(
                ward=ward, year=int(year), month=int(month), schedule_type=schedule_type
            ).first()

            month_prefix = f"{year}-{str(month).zfill(2)}"
            has_final = FinalSchedule.objects.filter(
                ward=ward,
                schedule_date__startswith=month_prefix,
                schedule_type=schedule_type
            ).exists()

            has_draft = DraftSchedule.objects.filter(
                ward=ward,
                schedule_date__startswith=month_prefix,
                source='manual',
                schedule_type=schedule_type
            ).exists()

            from ai_module.models import AISchedule, SpecialScheduleConfig
            has_ai = AISchedule.objects.filter(
                ward=ward, year=int(year), month=int(month), schedule_type=schedule_type
            ).exists()

            has_config = True
            # ✅ شامل assistant
            if schedule_type in ('oncall', 'crisis', 'assistant'):
                has_config = SpecialScheduleConfig.objects.filter(
                    ward=ward, schedule_type=schedule_type
                ).exists()

            # ✅ فراخوانی can_generate
            can_generate, error_msg, missing_prereq = WorkflowEngine.can_generate(
                ward, int(year), int(month), schedule_type
            )

            # ✅ وضعیت واقعی
            actual_status = WorkflowEngine.get_status(
                ward, int(year), int(month), schedule_type
            )

            # ✅ نمایش فارسی وضعیت
            status_display = {
                'DRAFT': '⏳ پیش‌نویس',
                'PENDING_SUPERVISOR': '⏳ در انتظار تایید سوپروایزر',
                'PENDING_MATRON': '⏳ در انتظار تایید نهایی مترون',
                'LOCKED': '🔒 قفل شده',
                'RETURNED_TO_HEAD': '🔄 بازگشته به سرپرستار',
                'RETURNED_TO_SUPERVISOR': '🔄 بازگشته به سوپروایزر',
            }.get(actual_status, 'نامشخص')

            result[schedule_type] = {
                'status': actual_status,
                'status_display': status_display,
                'has_final': has_final,
                'has_draft': has_draft,
                'has_ai': has_ai,
                'has_config': has_config,
                'can_generate': can_generate,
                'generation_error': error_msg,
                'missing_prerequisite': missing_prereq,
                'ward_notes': status_obj.ward_notes if status_obj else '',
                'supervisor_feedback': status_obj.supervisor_feedback if status_obj else '',
                'matron_feedback': status_obj.matron_feedback if status_obj else '',
                'updated_at': status_obj.updated_at.isoformat() if status_obj else None,
            }

        return Response(result, status=status.HTTP_200_OK)


# ========================================================
# ✅ توابع کمکی برای Excel/PDF
# ========================================================
def _get_used_shift_codes_and_shifts(ward, year, month, schedule_type):
    """
    خواندن shiftها از FinalSchedule + کدهای استفاده‌شده.
    خروجی: (shifts_dict, used_shift_codes_qs)
    """
    final_schedules = FinalSchedule.objects.filter(
        ward=ward,
        schedule_date__startswith=f"{year}-{str(month).zfill(2)}",
        schedule_type=schedule_type
    ).select_related('assigned_staff', 'shift_code')

    shifts_dict = {}
    used_shift_codes_set = set()

    for fs in final_schedules:
        staff_id = fs.assigned_staff.id
        day = int(fs.schedule_date.split('-')[2])
        if staff_id not in shifts_dict:
            shifts_dict[staff_id] = {}
        if fs.shift_code:
            shifts_dict[staff_id][day] = fs.shift_code.code
            used_shift_codes_set.add(fs.shift_code.code)

    used_shift_codes = ShiftCode.objects.filter(
        code__in=used_shift_codes_set
    ).order_by('code')

    return shifts_dict, used_shift_codes


def _get_export_schedule_type_label(schedule_type):
    """برچسب کامل برای فایل/صفحه"""
    return {
        'main': 'برنامه اصلی',
        'oncall': 'برنامه آنکالی',
        'crisis': 'برنامه بحران',
        'assistant': 'برنامه کمک پرستار',
    }.get(schedule_type, 'برنامه')


def _get_export_type_colors(schedule_type):
    """رنگ‌های badge برای هر نوع"""
    colors = {
        'main': {'bg': 'DBEAFE', 'text': '1E40AF'},
        'oncall': {'bg': 'FEF3C7', 'text': '92400E'},
        'crisis': {'bg': 'FEE2E2', 'text': '991B1B'},
        'assistant': {'bg': 'D1FAE5', 'text': '065F46'},
    }
    return colors.get(schedule_type, colors['main'])


def _get_export_type_colors_html(schedule_type):
    """رنگ‌های badge HTML برای PDF"""
    colors = {
        'main': {'bg': '#DBEAFE', 'text': '#1E40AF'},
        'oncall': {'bg': '#FEF3C7', 'text': '#92400E'},
        'crisis': {'bg': '#FEE2E2', 'text': '#991B1B'},
        'assistant': {'bg': '#D1FAE5', 'text': '#065F46'},
    }
    return colors.get(schedule_type, colors['main'])


def _add_schedule_sheet_to_excel(wb, ward, year, month, schedule_type, site_settings):
    """
    ✅ اضافه کردن یک شیت به workbook برای یک نوع برنامه.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter

    month_prefix = f"{year}-{str(month).zfill(2)}"

    shifts_dict, used_shift_codes = _get_used_shift_codes_and_shifts(
        ward, year, month, schedule_type
    )

    profiles = list(_get_profiles_for_schedule_type(ward, schedule_type, year, month))
    profiles = sorted(profiles, key=lambda p: (
        int(p.position) if p.position and p.position.isdigit() else 999,
        -p.work_experience
    ))

    # ✅ نکات
    ward_notes_text = ''
    monthly_note = WardMonthlyNotes.objects.filter(
        ward=ward, year=int(year), month=int(month)
    ).first()
    if monthly_note and monthly_note.notes:
        ward_notes_text = monthly_note.notes

    # ✅ روزها
    days_in_month = get_persian_days_in_month(int(year), int(month))
    month_names = ['', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور', 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند']
    weekdays = ['شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه']
    first_day_index = get_persian_weekday(int(year), int(month), 1)

    weekday_short = {
        'شنبه': 'ش', 'یکشنبه': 'ی', 'دوشنبه': 'د',
        'سه‌شنبه': 'س', 'چهارشنبه': 'چ', 'پنجشنبه': 'پ', 'جمعه': 'ج',
    }

    holidays = Holiday.objects.filter(
        date__startswith=month_prefix
    ).values('date', 'name')
    holidays_dict = {h['date']: h['name'] for h in holidays}

    schedule_type_label = _get_export_schedule_type_label(schedule_type)
    type_color = _get_export_type_colors(schedule_type)

    # ✅ targetها
    (
        target_m, target_e, target_n,
        holiday_target_m, holiday_target_e, holiday_target_n
    ) = _get_targets_for_schedule_type(ward, schedule_type)

    # ============================================================
    # ✅ ساخت شیت
    # ============================================================
    ws = wb.create_sheet(title=schedule_type_label)
    ws.sheet_view.rightToLeft = True

    title_font = Font(name='B Yagut', size=14, bold=True, color='2C3E50')
    subtitle_font = Font(name='B Yagut', size=11, bold=True, color='34495E')
    header_font = Font(name='B Yagut', size=10, bold=True, color='FFFFFF')
    content_font = Font(name='B Zar', size=9)
    bold_font = Font(name='B Zar', size=9, bold=True)
    type_badge_font = Font(name='B Yagut', size=11, bold=True, color=type_color['text'])

    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    type_fill = PatternFill(start_color=type_color['bg'], end_color=type_color['bg'], fill_type='solid')
    stat_m_fill = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid')
    stat_e_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    stat_n_fill = PatternFill(start_color='E0E7FF', end_color='E0E7FF', fill_type='solid')
    notes_fill = PatternFill(start_color='FFF8E1', end_color='FFF8E1', fill_type='solid')
    codes_fill = PatternFill(start_color='F0FDF4', end_color='F0FDF4', fill_type='solid')
    alternate_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')

    shift_fills = {
        'M': PatternFill(start_color='3498DB', end_color='3498DB', fill_type='solid'),
        'E': PatternFill(start_color='F39C12', end_color='F39C12', fill_type='solid'),
        'N': PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid'),
        'ME': PatternFill(start_color='9B59B6', end_color='9B59B6', fill_type='solid'),
        'MN': PatternFill(start_color='16A085', end_color='16A085', fill_type='solid'),
        'EN': PatternFill(start_color='E67E22', end_color='E67E22', fill_type='solid'),
        'MEN': PatternFill(start_color='8E44AD', end_color='8E44AD', fill_type='solid'),
        'H': PatternFill(start_color='E74C3C', end_color='E74C3C', fill_type='solid'),
        '-': PatternFill(start_color='95A5A6', end_color='95A5A6', fill_type='solid'),
    }

    shift_fonts = {
        'M': Font(name='B Zar', size=9, bold=True, color='FFFFFF'),
        'E': Font(name='B Zar', size=9, bold=True, color='FFFFFF'),
        'N': Font(name='B Zar', size=9, bold=True, color='FFFFFF'),
        'ME': Font(name='B Zar', size=8, bold=True, color='FFFFFF'),
        'MN': Font(name='B Zar', size=8, bold=True, color='FFFFFF'),
        'EN': Font(name='B Zar', size=8, bold=True, color='FFFFFF'),
        'MEN': Font(name='B Zar', size=7, bold=True, color='FFFFFF'),
        'H': Font(name='B Zar', size=9, bold=True, color='FFFFFF'),
        '-': Font(name='B Zar', size=9, bold=True, color='FFFFFF'),
    }

    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    medium_border = Border(
        left=Side(style='medium', color='2C3E50'),
        right=Side(style='medium', color='2C3E50'),
        top=Side(style='medium', color='2C3E50'),
        bottom=Side(style='medium', color='2C3E50')
    )

    total_cols = 4 + days_in_month + 5
    last_col_letter = get_column_letter(total_cols)

    # لوگوها
    if site_settings and site_settings.university_logo:
        try:
            img = XLImage(site_settings.university_logo.path)
            img.width = 55
            img.height = 55
            ws.add_image(img, 'A1')
        except Exception:
            pass

    if site_settings and site_settings.hospital_logo:
        try:
            img = XLImage(site_settings.hospital_logo.path)
            img.width = 55
            img.height = 55
            ws.add_image(img, f'{last_col_letter}1')
        except Exception:
            pass

    # هدر
    ws.merge_cells(f'A1:{last_col_letter}1')
    ws['A1'] = site_settings.hospital_name if site_settings else ''
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'A2:{last_col_letter}2')
    ws['A2'] = f"{ward.name} - {month_names[int(month)]} {year}"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'A3:{last_col_letter}3')
    type_cell = ws['A3']
    type_cell.value = f"📋 {schedule_type_label}"
    type_cell.font = type_badge_font
    type_cell.fill = type_fill
    type_cell.alignment = Alignment(horizontal='center', vertical='center')
    type_cell.border = medium_border

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 22

    # هدر جدول
    headers = ['ردیف', 'نام و نام خانوادگی', 'سمت', 'وضعیت استخدام']
    for day in range(1, days_in_month + 1):
        day_index = first_day_index + (day - 1)
        weekday_full = weekdays[day_index % 7]
        weekday_short_char = weekday_short.get(weekday_full, weekday_full[0])
        headers.append(f"{day}\n{weekday_short_char}")
    headers.extend(['صبح', 'عصر', 'شب', 'مرخصی', 'آف'])

    header_row = 5
    ws.row_dimensions[header_row].height = 30

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    # داده‌ها
    data_row = header_row + 1
    for idx, profile in enumerate(profiles):
        row_data = [
            idx + 1,
            f"{profile.user.first_name} {profile.user.last_name}",
            profile.get_position_display(),
            profile.employment_status
        ]

        total_m = total_e = total_n = total_h = total_off = 0

        for day in range(1, days_in_month + 1):
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            row_data.append(shift_code if shift_code else '-')
            if shift_code:
                if 'M' in shift_code: total_m += 1
                if 'E' in shift_code: total_e += 1
                if 'N' in shift_code: total_n += 1
                if 'H' in shift_code: total_h += 1
                if '-' in shift_code: total_off += 1

        row_data.extend([total_m, total_e, total_n, total_h, total_off])

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=data_row, column=col)
            cell.value = value
            cell.font = content_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

            if 5 <= col <= 4 + days_in_month:
                shift_code = value
                if shift_code in shift_fills:
                    cell.fill = shift_fills[shift_code]
                    cell.font = shift_fonts[shift_code]
            else:
                if idx % 2 == 1:
                    cell.fill = alternate_fill

            if col == 2:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal='right', vertical='center')

        data_row += 1

    # ردیف آماری صبح
    stat_m_row = data_row
    ws.row_dimensions[stat_m_row].height = 20

    ws.merge_cells(f'A{stat_m_row}:D{stat_m_row}')
    label_cell = ws.cell(row=stat_m_row, column=1)
    label_cell.value = f"☀️ صبح (M) عادی: {target_m} / تعطیل: {holiday_target_m}"
    label_cell.font = bold_font
    label_cell.fill = stat_m_fill
    label_cell.alignment = Alignment(horizontal='center', vertical='center')
    label_cell.border = thin_border

    for day in range(1, days_in_month + 1):
        count_m = 0
        day_index = first_day_index + (day - 1)
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)
        target = holiday_target_m if is_holiday_or_friday else target_m

        for profile in profiles:
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code and 'M' in shift_code:
                count_m += 1

        cell = ws.cell(row=stat_m_row, column=4 + day)
        cell.value = count_m
        cell.font = bold_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

        if count_m < target:
            cell.fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
        elif count_m > target:
            cell.fill = PatternFill(start_color='FFEDD5', end_color='FFEDD5', fill_type='solid')
        else:
            cell.fill = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')

    total_m_all = sum(sum(1 for d in shifts_dict.get(p.user.id, {}).values() if 'M' in d) for p in profiles)
    total_cell = ws.cell(row=stat_m_row, column=4 + days_in_month + 1)
    total_cell.value = total_m_all
    total_cell.font = bold_font
    total_cell.fill = stat_m_fill
    total_cell.alignment = Alignment(horizontal='center', vertical='center')
    total_cell.border = thin_border

    # ردیف آماری عصر
    stat_e_row = stat_m_row + 1
    ws.row_dimensions[stat_e_row].height = 20

    ws.merge_cells(f'A{stat_e_row}:D{stat_e_row}')
    label_cell = ws.cell(row=stat_e_row, column=1)
    label_cell.value = f"🌅 عصر (E) عادی: {target_e} / تعطیل: {holiday_target_e}"
    label_cell.font = bold_font
    label_cell.fill = stat_e_fill
    label_cell.alignment = Alignment(horizontal='center', vertical='center')
    label_cell.border = thin_border

    for day in range(1, days_in_month + 1):
        count_e = 0
        day_index = first_day_index + (day - 1)
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)
        target = holiday_target_e if is_holiday_or_friday else target_e

        for profile in profiles:
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code and 'E' in shift_code:
                count_e += 1

        cell = ws.cell(row=stat_e_row, column=4 + day)
        cell.value = count_e
        cell.font = bold_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

        if count_e < target:
            cell.fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
        elif count_e > target:
            cell.fill = PatternFill(start_color='FFEDD5', end_color='FFEDD5', fill_type='solid')
        else:
            cell.fill = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')

    total_e_all = sum(sum(1 for d in shifts_dict.get(p.user.id, {}).values() if 'E' in d) for p in profiles)
    total_cell = ws.cell(row=stat_e_row, column=4 + days_in_month + 1)
    total_cell.value = total_e_all
    total_cell.font = bold_font
    total_cell.fill = stat_e_fill
    total_cell.alignment = Alignment(horizontal='center', vertical='center')
    total_cell.border = thin_border

    # ردیف آماری شب
    stat_n_row = stat_e_row + 1
    ws.row_dimensions[stat_n_row].height = 20

    ws.merge_cells(f'A{stat_n_row}:D{stat_n_row}')
    label_cell = ws.cell(row=stat_n_row, column=1)
    label_cell.value = f"🌙 شب (N) عادی: {target_n} / تعطیل: {holiday_target_n}"
    label_cell.font = bold_font
    label_cell.fill = stat_n_fill
    label_cell.alignment = Alignment(horizontal='center', vertical='center')
    label_cell.border = thin_border

    for day in range(1, days_in_month + 1):
        count_n = 0
        day_index = first_day_index + (day - 1)
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)
        target = holiday_target_n if is_holiday_or_friday else target_n

        for profile in profiles:
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code and 'N' in shift_code:
                count_n += 1

        cell = ws.cell(row=stat_n_row, column=4 + day)
        cell.value = count_n
        cell.font = bold_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

        if count_n < target:
            cell.fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
        elif count_n > target:
            cell.fill = PatternFill(start_color='FFEDD5', end_color='FFEDD5', fill_type='solid')
        else:
            cell.fill = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')

    total_n_all = sum(sum(1 for d in shifts_dict.get(p.user.id, {}).values() if 'N' in d) for p in profiles)
    total_cell = ws.cell(row=stat_n_row, column=4 + days_in_month + 1)
    total_cell.value = total_n_all
    total_cell.font = bold_font
    total_cell.fill = stat_n_fill
    total_cell.alignment = Alignment(horizontal='center', vertical='center')
    total_cell.border = thin_border

    # کدهای استفاده‌شده
    codes_row = stat_n_row + 2
    ws.merge_cells(f'A{codes_row}:D{codes_row}')
    codes_label = ws.cell(row=codes_row, column=1)
    codes_label.value = '📋 کدهای شیفت استفاده‌شده در این برنامه:'
    codes_label.font = bold_font
    codes_label.fill = codes_fill
    codes_label.alignment = Alignment(horizontal='center', vertical='center')
    codes_label.border = thin_border

    codes_text_parts = [f"{sc.code} ({sc.label})" for sc in used_shift_codes]
    codes_text = ' | '.join(codes_text_parts) if codes_text_parts else 'هیچ کدی استفاده نشده است'

    ws.merge_cells(f'E{codes_row}:{last_col_letter}{codes_row}')
    codes_value = ws.cell(row=codes_row, column=5)
    codes_value.value = codes_text
    codes_value.font = content_font
    codes_value.fill = codes_fill
    codes_value.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    codes_value.border = thin_border

    # نکات مهم بخش
    notes_row = codes_row + 1
    ws.merge_cells(f'A{notes_row}:D{notes_row}')
    notes_label = ws.cell(row=notes_row, column=1)
    notes_label.value = '📝 نکات مهم بخش:'
    notes_label.font = bold_font
    notes_label.fill = notes_fill
    notes_label.alignment = Alignment(horizontal='center', vertical='center')
    notes_label.border = thin_border

    ws.merge_cells(f'E{notes_row}:{last_col_letter}{notes_row}')
    notes_value = ws.cell(row=notes_row, column=5)
    notes_value.value = ward_notes_text or 'نکته‌ای ثبت نشده است'
    notes_value.font = content_font
    notes_value.fill = notes_fill
    notes_value.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    notes_value.border = thin_border

    # امضاها
    signature_row = notes_row + 3
    ws.row_dimensions[signature_row].height = 60

    mid_col = max(1, total_cols // 2)
    mid_col_letter = get_column_letter(mid_col)

    ws.merge_cells(f'A{signature_row}:{mid_col_letter}{signature_row}')
    head_sig_cell = ws.cell(row=signature_row, column=1)
    head_sig_cell.value = '✍️ سرپرستار:'
    head_sig_cell.font = bold_font
    head_sig_cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'{get_column_letter(mid_col + 1)}{signature_row}:{last_col_letter}{signature_row}')
    matron_sig_cell = ws.cell(row=signature_row, column=mid_col + 1)
    matron_sig_cell.value = '✍️ مدیر خدمات پرستاری:'
    matron_sig_cell.font = bold_font
    matron_sig_cell.alignment = Alignment(horizontal='center', vertical='center')

    # عرض ستون‌ها
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 10
    for day in range(1, days_in_month + 1):
        ws.column_dimensions[get_column_letter(4 + day)].width = 5
    for i in range(5):
        ws.column_dimensions[get_column_letter(4 + days_in_month + 1 + i)].width = 7


# ========================================================
# خروجی اکسل
# ========================================================
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
import io


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_excel(request):
    """
    ✅ خروجی اکسل — برای main، دو شیت (اصلی + کمک پرستار) می‌سازد.
    برای بقیه، یک شیت.
    """
    try:
        ward_id = request.query_params.get('ward_id')
        year = request.query_params.get('year')
        month = request.query_params.get('month')
        schedule_type = request.query_params.get('schedule_type', 'main')

        if not ward_id or not year or not month:
            return Response({'error': 'شناسه بخش، سال و ماه الزامی است.'}, status=400)

        try:
            ward = Ward.objects.get(id=ward_id)
        except Ward.DoesNotExist:
            return Response({'error': 'بخش مورد نظر یافت نشد.'}, status=404)

        site_settings = SiteSettings.objects.filter(is_active=True).first()
        month_prefix = f"{year}-{str(month).zfill(2)}"

        # ============================================================
        # ✅ چک قفل بودن برنامه
        # ============================================================
        is_locked = FinalSchedule.objects.filter(
            ward=ward,
            schedule_date__startswith=month_prefix,
            is_locked=True,
            schedule_type=schedule_type
        ).exists()

        if not is_locked:
            return Response({'error': 'خروجی فقط بعد از تایید نهایی مترون و قفل شدن برنامه قابل دریافت است.'}, status=403)

        # ============================================================
        # ✅ ساخت workbook
        # ============================================================
        wb = Workbook()
        wb.remove(wb.active)   # حذف شیت پیش‌فرض

        # ✅ برای main، دو شیت (اصلی + کمک پرستار)
        if schedule_type == 'main':
            schedule_types_to_export = ['main']
            # چک کن که assistant هم قفل شده باشه
            is_assistant_locked = FinalSchedule.objects.filter(
                ward=ward,
                schedule_date__startswith=month_prefix,
                is_locked=True,
                schedule_type='assistant'
            ).exists()
            if is_assistant_locked:
                schedule_types_to_export.append('assistant')
        else:
            schedule_types_to_export = [schedule_type]

        for st in schedule_types_to_export:
            _add_schedule_sheet_to_excel(
                wb=wb,
                ward=ward,
                year=int(year),
                month=int(month),
                schedule_type=st,
                site_settings=site_settings,
            )

        # ✅ اگه هیچ شیتی ساخته نشد
        if len(wb.sheetnames) == 0:
            return Response({'error': 'هیچ برنامه‌ای برای خروجی یافت نشد.'}, status=404)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        # ✅ نام فایل
        schedule_type_label = _get_export_schedule_type_label(schedule_type).replace(' ', '_')

        month_names = ['', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور', 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند']
        filename = f"{schedule_type_label}_{ward.name}_{month_names[int(month)]}_{year}.xlsx"

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
        return response

    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': f'خطا در خروجی اکسل: {str(e)}'}, status=500)


# ========================================================
# خروجی PDF
# ========================================================
from weasyprint import HTML
import os
from django.conf import settings


def get_asset_path(path):
    if path:
        return f"file:///{path}"
    return ''


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_pdf(request):
    """
    ✅ خروجی PDF — برای main، دو صفحه (اصلی + کمک پرستار) می‌سازد.
    برای بقیه، یک صفحه.
    """
    import base64
    import traceback

    try:
        ward_id = request.query_params.get('ward_id')
        year = request.query_params.get('year')
        month = request.query_params.get('month')
        schedule_type = request.query_params.get('schedule_type', 'main')

        if not ward_id or not year or not month:
            return Response({'error': 'شناسه بخش، سال و ماه الزامی است.'}, status=400)

        try:
            ward = Ward.objects.get(id=ward_id)
        except Ward.DoesNotExist:
            return Response({'error': 'بخش مورد نظر یافت نشد.'}, status=404)

        site_settings = SiteSettings.objects.filter(is_active=True).first()
        month_prefix = f"{year}-{str(month).zfill(2)}"

        # ============================================================
        # ✅ چک قفل بودن برنامه
        # ============================================================
        is_locked = FinalSchedule.objects.filter(
            ward=ward,
            schedule_date__startswith=month_prefix,
            is_locked=True,
            schedule_type=schedule_type
        ).exists()

        if not is_locked:
            return Response({'error': 'خروجی فقط بعد از تایید نهایی مترون و قفل شدن برنامه قابل دریافت است.'}, status=403)

        # ✅ برای main، دو صفحه
        if schedule_type == 'main':
            schedule_types_to_export = ['main']
            is_assistant_locked = FinalSchedule.objects.filter(
                ward=ward,
                schedule_date__startswith=month_prefix,
                is_locked=True,
                schedule_type='assistant'
            ).exists()
            if is_assistant_locked:
                schedule_types_to_export.append('assistant')
        else:
            schedule_types_to_export = [schedule_type]

        # ============================================================
        # ✅ ساخت HTML کلی
        # ============================================================
        html_parts = []
        for st in schedule_types_to_export:
            html_parts.append(
                _build_pdf_page_for_schedule(
                    ward=ward,
                    year=int(year),
                    month=int(month),
                    schedule_type=st,
                    site_settings=site_settings,
                )
            )

        html_content = _build_pdf_document(html_parts)

        from weasyprint import HTML
        pdf_file = HTML(string=html_content, base_url=settings.BASE_DIR).write_pdf()

        schedule_type_label = _get_export_schedule_type_label(schedule_type).replace(' ', '_')
        month_names = ['', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور', 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند']
        filename = f"{schedule_type_label}_{ward.name}_{month_names[int(month)]}_{year}.pdf"

        response = HttpResponse(pdf_file, content_type='application/pdf')
        response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
        return response

    except Exception as e:
        traceback.print_exc()
        return Response({'error': f'خطا در خروجی PDF: {str(e)}'}, status=500)


def _build_pdf_document(html_parts):
    """ساخت HTML کلی با چند صفحه"""
    css = """
    <style>
        @page {
            size: A4 landscape;
            margin: 4mm;
        }

        * { box-sizing: border-box; }

        body {
            font-family: 'Vazirmatn', 'B Nazanin', sans-serif;
            font-size: 11px;
            margin: 0;
            padding: 0;
            color: #333;
            direction: rtl;
        }

        .page {
            page-break-after: always;
        }

        .page:last-child {
            page-break-after: auto;
        }

        .header {
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 5px;
            border-bottom: 3px solid #2c3e50;
            padding-bottom: 4px;
        }

        .logo {
            width: 50px;
            height: 50px;
            object-fit: contain;
        }

        .header-center {
            text-align: center;
            flex: 1;
            padding: 0 8px;
        }

        .university-name {
            font-size: 15px;
            font-weight: bold;
            color: #2c3e50;
            margin: 0;
        }

        .hospital-name {
            font-size: 13px;
            font-weight: bold;
            color: #34495e;
            margin: 2px 0;
        }

        .ward-name {
            font-size: 12px;
            color: #7f8c8d;
            margin: 0;
        }

        .type-badge {
            display: inline-block;
            padding: 3px 16px;
            border-radius: 10px;
            font-weight: bold;
            font-size: 12px;
            margin: 4px 0;
        }

        table {
            width: 100%;
            border-collapse: collapse;
            margin-top: 3px;
            table-layout: fixed;
        }

        th {
            background-color: #2c3e50;
            color: white;
            padding: 5px 1px;
            font-size: 10px;
            font-weight: bold;
            border: 1px solid #34495e;
            text-align: center;
            white-space: nowrap;
            height: 38px;
            overflow: hidden;
        }

        td {
            padding: 4px 0px;
            border: 1px solid #ddd;
            text-align: center;
            font-size: 11px;
            font-weight: bold;
            white-space: nowrap;
            overflow: hidden;
            height: 26px;
        }

        .col-index { width: 22px; }
        .col-name { width: 90px; }
        .col-position { width: 55px; }
        .col-employment { width: 50px; }
        .col-day { width: 24px; }
        .col-stat { width: 26px; }

        .staff-name {
            text-align: right;
            font-weight: bold;
            font-size: 11px;
            padding-right: 4px;
            overflow: hidden;
            text-overflow: ellipsis;
        }

        .stat-label {
            text-align: center;
            font-weight: bold;
            font-size: 10px;
        }

        .stat-m { background-color: #DBEAFE; }
        .stat-e { background-color: #FEF3C7; }
        .stat-n { background-color: #E0E7FF; }

        .shift-M { background-color: #3498DB; color: white; font-weight: bold; font-size: 12px; }
        .shift-E { background-color: #F39C12; color: white; font-weight: bold; font-size: 12px; }
        .shift-N { background-color: #2C3E50; color: white; font-weight: bold; font-size: 12px; }
        .shift-ME { background-color: #9B59B6; color: white; font-weight: bold; font-size: 9px; }
        .shift-MN { background-color: #16A085; color: white; font-weight: bold; font-size: 9px; }
        .shift-EN { background-color: #E67E22; color: white; font-weight: bold; font-size: 9px; }
        .shift-MEN { background-color: #8E44AD; color: white; font-weight: bold; font-size: 8px; }
        .shift-H { background-color: #E74C3C; color: white; font-weight: bold; font-size: 12px; }
        .shift-off { background-color: #95A5A6; color: white; font-weight: bold; font-size: 12px; }

        .day-number {
            font-size: 13px;
            font-weight: 900;
            line-height: 1.1;
        }

        .day-weekday {
            font-size: 10px;
            font-weight: 700;
            opacity: 0.95;
            line-height: 1.1;
        }

        .codes-section {
            margin-top: 10px;
            padding: 8px 12px;
            background-color: #F0FDF4;
            border-radius: 4px;
            border-right: 5px solid #10B981;
        }

        .codes-content {
            font-size: 12px;
            color: #065F46;
            font-weight: bold;
        }

        .notes-section {
            margin-top: 10px;
            padding: 8px 12px;
            background-color: #FFF8E1;
            border-radius: 4px;
            border-right: 5px solid #F39C12;
        }

        .notes-content {
            font-size: 12px;
            line-height: 1.8;
            color: #2c3e50;
        }

        .signatures {
            display: flex;
            justify-content: space-around;
            margin-top: 25px;
            padding-top: 15px;
            border-top: 2px solid #ddd;
        }

        .signature-box {
            text-align: center;
            width: 240px;
        }

        .signature-title {
            font-size: 13px;
            font-weight: bold;
            color: #2c3e50;
            margin-bottom: 8px;
        }

        .signature-img {
            max-height: 70px;
            max-width: 170px;
            object-fit: contain;
        }

        .signature-empty {
            color: #95a5a6;
            font-size: 11px;
            font-style: italic;
            padding: 25px 0;
        }
    </style>
    """

    html = f"""<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
    <meta charset="UTF-8">
    {css}
</head>
<body>
    {''.join(html_parts)}
</body>
</html>"""

    return html


def _build_pdf_page_for_schedule(ward, year, month, schedule_type, site_settings):
    """
    ✅ ساخت HTML یک صفحه برای یک نوع برنامه.
    """
    import base64
    from django.conf import settings

    month_prefix = f"{year}-{str(month).zfill(2)}"

    shifts_dict, used_shift_codes = _get_used_shift_codes_and_shifts(
        ward, year, month, schedule_type
    )

    profiles = list(_get_profiles_for_schedule_type(ward, schedule_type, year, month))
    profiles = sorted(profiles, key=lambda p: (
        int(p.position) if p.position and p.position.isdigit() else 999,
        -p.work_experience
    ))

    ward_notes_text = ''
    monthly_note = WardMonthlyNotes.objects.filter(
        ward=ward, year=int(year), month=int(month)
    ).first()
    if monthly_note and monthly_note.notes:
        ward_notes_text = monthly_note.notes

    days_in_month = get_persian_days_in_month(int(year), int(month))
    month_names = ['', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور', 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند']
    weekdays = ['شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه']
    first_day_index = get_persian_weekday(int(year), int(month), 1)

    weekday_short = {
        'شنبه': 'ش', 'یکشنبه': 'ی', 'دوشنبه': 'د',
        'سه‌شنبه': 'س', 'چهارشنبه': 'چ', 'پنجشنبه': 'پ', 'جمعه': 'ج',
    }

    holidays = Holiday.objects.filter(
        date__startswith=month_prefix
    ).values('date', 'name')
    holidays_dict = {h['date']: h['name'] for h in holidays}

    schedule_type_label = _get_export_schedule_type_label(schedule_type)
    tc = _get_export_type_colors_html(schedule_type)

    (
        target_m, target_e, target_n,
        holiday_target_m, holiday_target_e, holiday_target_n
    ) = _get_targets_for_schedule_type(ward, schedule_type)

    def get_image_base64(image_field):
        if not image_field:
            return ''
        try:
            with open(image_field.path, 'rb') as f:
                data = f.read()
            ext = image_field.path.split('.')[-1].lower()
            mime = 'image/png' if ext == 'png' else (
                'image/jpeg' if ext in ['jpg', 'jpeg'] else 'image/png'
            )
            b64 = base64.b64encode(data).decode('utf-8')
            return f"data:{mime};base64,{b64}"
        except Exception as e:
            print(f"⚠️ خطا در base64: {e}")
            return ''

    university_logo_url = get_image_base64(site_settings.university_logo) if site_settings else ''
    hospital_logo_url = get_image_base64(site_settings.hospital_logo) if site_settings else ''

    head_profile = Profile.objects.filter(ward=ward, role='head').first()
    matron_profile = Profile.objects.filter(role='admin').first()

    head_sig = get_image_base64(head_profile.signature) if head_profile and head_profile.signature else ''
    matron_sig = get_image_base64(matron_profile.signature) if matron_profile and matron_profile.signature else ''

    # ============================================================
    # ✅ ساخت HTML صفحه
    # ============================================================
    html = f"""
    <div class="page">
        <div class="header">
            <div class="header-left">
                {f'<img src="{university_logo_url}" class="logo" />' if university_logo_url else ''}
            </div>
            <div class="header-center">
                <h1 class="university-name">{site_settings.university_name if site_settings else ''}</h1>
                <h2 class="hospital-name">{site_settings.hospital_name if site_settings else ''}</h2>
                <p class="ward-name">{ward.name} - {month_names[int(month)]} {year}</p>
                <div class="type-badge" style="background: {tc['bg']}; color: {tc['text']}; border: 2px solid {tc['text']};">📋 {schedule_type_label}</div>
            </div>
            <div class="header-right">
                {f'<img src="{hospital_logo_url}" class="logo" />' if hospital_logo_url else ''}
            </div>
        </div>

        <table>
            <colgroup>
                <col class="col-index" />
                <col class="col-name" />
                <col class="col-position" />
                <col class="col-employment" />
    """

    for day in range(1, days_in_month + 1):
        html += '<col class="col-day" />'

    for _ in range(5):
        html += '<col class="col-stat" />'

    html += """
            </colgroup>
            <thead>
                <tr>
                    <th>ردیف</th>
                    <th>نام و نام خانوادگی</th>
                    <th>سمت</th>
                    <th>وضعیت استخدام</th>
    """

    for day in range(1, days_in_month + 1):
        day_index = first_day_index + (day - 1)
        weekday_full = weekdays[day_index % 7]
        weekday_short_char = weekday_short.get(weekday_full, weekday_full[0])
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday = date_key in holidays_dict
        is_friday = (day_index % 7 == 6)

        bg = ''
        if is_holiday:
            bg = 'background-color: #FEE2E2;'
        elif is_friday:
            bg = 'background-color: #FEF3C7;'

        html += f'''
                    <th style="{bg}">
                        <div class="day-number">{day}</div>
                        <div class="day-weekday">{weekday_short_char}</div>
                    </th>
        '''

    html += '''
                    <th style="background: #DBEAFE; color: #1E40AF; font-size: 11px;">صبح</th>
                    <th style="background: #FEF3C7; color: #92400E; font-size: 11px;">عصر</th>
                    <th style="background: #E0E7FF; color: #3730A3; font-size: 11px;">شب</th>
                    <th style="background: #FEE2E2; color: #991B1B; font-size: 10px;">مرخصی</th>
                    <th style="background: #F3F4F6; color: #374151; font-size: 11px;">آف</th>
                </tr>
            </thead>
            <tbody>
    '''

    for idx, profile in enumerate(profiles):
        total_m = total_e = total_n = total_h = total_off = 0
        row_bg = '#F8F9FA' if idx % 2 == 1 else 'white'

        html += f'<tr style="background-color: {row_bg};">'
        html += f'<td style="font-weight: bold; color: #6B7280; font-size: 11px;">{idx + 1}</td>'
        html += f'<td class="staff-name">{profile.user.first_name} {profile.user.last_name}</td>'
        html += f'<td style="font-size: 10px;">{profile.get_position_display()}</td>'
        html += f'<td style="font-size: 10px;">{profile.employment_status}</td>'

        for day in range(1, days_in_month + 1):
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code:
                shift_class = f'shift-{shift_code}' if shift_code != '-' else 'shift-off'
                html += f'<td class="{shift_class}">{shift_code}</td>'
                if 'M' in shift_code: total_m += 1
                if 'E' in shift_code: total_e += 1
                if 'N' in shift_code: total_n += 1
                if 'H' in shift_code: total_h += 1
                if '-' in shift_code: total_off += 1
            else:
                html += '<td style="color: #ccc;">-</td>'

        html += f'''
            <td style="background: #DBEAFE; font-weight: bold; font-size: 11px;">{total_m}</td>
            <td style="background: #FEF3C7; font-weight: bold; font-size: 11px;">{total_e}</td>
            <td style="background: #E0E7FF; font-weight: bold; font-size: 11px;">{total_n}</td>
            <td style="background: #FEE2E2; font-weight: bold; font-size: 11px;">{total_h}</td>
            <td style="background: #F3F4F6; font-weight: bold; font-size: 11px;">{total_off}</td>
        </tr>
        '''

    # آمار صبح
    html += f'''
        <tr class="stat-m">
            <td colspan="4" class="stat-label">☀️ صبح (M) عادی: {target_m} / تعطیل: {holiday_target_m}</td>
    '''
    for day in range(1, days_in_month + 1):
        count_m = 0
        day_index = first_day_index + (day - 1)
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)
        target = holiday_target_m if is_holiday_or_friday else target_m

        for profile in profiles:
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code and 'M' in shift_code:
                count_m += 1

        bg_color = '#D1FAE5'
        if count_m < target:
            bg_color = '#FEE2E2'
        elif count_m > target:
            bg_color = '#FFEDD5'

        html += f'<td style="background: {bg_color}; font-weight: bold; font-size: 11px;">{count_m}</td>'

    total_m_all = sum(sum(1 for d in shifts_dict.get(p.user.id, {}).values() if 'M' in d) for p in profiles)
    html += f'<td style="background: #DBEAFE; font-weight: bold; font-size: 11px;">{total_m_all}</td>'
    html += '<td colspan="4"></td></tr>'

    # آمار عصر
    html += f'''
        <tr class="stat-e">
            <td colspan="4" class="stat-label">🌅 عصر (E) عادی: {target_e} / تعطیل: {holiday_target_e}</td>
    '''
    for day in range(1, days_in_month + 1):
        count_e = 0
        day_index = first_day_index + (day - 1)
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)
        target = holiday_target_e if is_holiday_or_friday else target_e

        for profile in profiles:
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code and 'E' in shift_code:
                count_e += 1

        bg_color = '#D1FAE5'
        if count_e < target:
            bg_color = '#FEE2E2'
        elif count_e > target:
            bg_color = '#FFEDD5'

        html += f'<td style="background: {bg_color}; font-weight: bold; font-size: 11px;">{count_e}</td>'

    total_e_all = sum(sum(1 for d in shifts_dict.get(p.user.id, {}).values() if 'E' in d) for p in profiles)
    html += f'<td style="background: #FEF3C7; font-weight: bold; font-size: 11px;">{total_e_all}</td>'
    html += '<td colspan="4"></td></tr>'

    # آمار شب
    html += f'''
        <tr class="stat-n">
            <td colspan="4" class="stat-label">🌙 شب (N) عادی: {target_n} / تعطیل: {holiday_target_n}</td>
    '''
    for day in range(1, days_in_month + 1):
        count_n = 0
        day_index = first_day_index + (day - 1)
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)
        target = holiday_target_n if is_holiday_or_friday else target_n

        for profile in profiles:
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code and 'N' in shift_code:
                count_n += 1

        bg_color = '#D1FAE5'
        if count_n < target:
            bg_color = '#FEE2E2'
        elif count_n > target:
            bg_color = '#FFEDD5'

        html += f'<td style="background: {bg_color}; font-weight: bold; font-size: 11px;">{count_n}</td>'

    total_n_all = sum(sum(1 for d in shifts_dict.get(p.user.id, {}).values() if 'N' in d) for p in profiles)
    html += f'<td style="background: #E0E7FF; font-weight: bold; font-size: 11px;">{total_n_all}</td>'
    html += '<td colspan="4"></td></tr>'

    html += '</tbody></table>'

    # کدهای استفاده‌شده
    codes_text = ' | '.join([f"{sc.code} ({sc.label})" for sc in used_shift_codes]) or 'هیچ کدی استفاده نشده است'
    html += f'''
    <div class="codes-section">
        <div style="font-size: 12px; font-weight: bold; color: #065F46; margin-bottom: 4px;">📋 کدهای شیفت استفاده‌شده در این برنامه:</div>
        <div class="codes-content">{codes_text}</div>
    </div>
    '''

    # نکات مهم
    html += f'''
    <div class="notes-section">
        <div style="font-size: 12px; font-weight: bold; color: #92400E; margin-bottom: 4px;">📝 نکات مهم بخش:</div>
        <div class="notes-content">{ward_notes_text or 'نکته‌ای ثبت نشده است'}</div>
    </div>
    '''

    # امضاها
    head_sig_html = f'<img src="{head_sig}" class="signature-img" />' if head_sig else '<div class="signature-empty">امضا ثبت نشده است</div>'
    matron_sig_html = f'<img src="{matron_sig}" class="signature-img" />' if matron_sig else '<div class="signature-empty">امضا ثبت نشده است</div>'

    html += f'''
    <div class="signatures">
        <div class="signature-box">
            <div class="signature-title">✍️ سرپرستار</div>
            {head_sig_html}
        </div>
        <div class="signature-box">
            <div class="signature-title">✍️ مدیر خدمات پرستاری</div>
            {matron_sig_html}
        </div>
    </div>
    </div>
    '''

    return html


# ========================================================
# ویو دریافت تنظیمات سراسری سایت
# ========================================================
@api_view(['GET'])
@permission_classes([AllowAny])
def get_site_settings(request):
    """
    دریافت تنظیمات سراسری سایت (نام بیمارستان، لوگوها، پاورقی و ...).
    """
    try:
        settings_obj = SiteSettings.objects.filter(is_active=True).first()
        if settings_obj:
            serializer = SiteSettingsSerializer(settings_obj, context={'request': request})
            return Response(serializer.data)

        return Response({
            'university_name': '',
            'university_logo': None,
            'hospital_name': '',
            'hospital_logo': None,
            'login_background': None,
            'footer_text': ''
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=500)


# ========================================================
# تحلیل هشدارها
# ========================================================
def analyze_schedule_warnings(shifts_dict, profiles, days_in_month, first_day_index,
                                holidays_dict, year, month, ward, schedule_type,
                                daily_target_m, daily_target_e, daily_target_n):
    """
    تحلیل هشدارهای یک برنامه.
    """
    warnings = []

    # ============================================================
    # 1. کمبود/اضافه پرسنل
    # ============================================================
    for day in range(1, days_in_month + 1):
        day_index = first_day_index + (day - 1)
        date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
        is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)

        count_m = 0
        count_e = 0
        count_n = 0

        for profile in profiles:
            shift_code = shifts_dict.get(profile.user.id, {}).get(day, '')
            if shift_code:
                if 'M' in shift_code: count_m += 1
                if 'E' in shift_code: count_e += 1
                if 'N' in shift_code: count_n += 1

        target_m = daily_target_m.get('holiday' if is_holiday_or_friday else 'normal', 0)
        target_e = daily_target_e.get('holiday' if is_holiday_or_friday else 'normal', 0)
        target_n = daily_target_n.get('holiday' if is_holiday_or_friday else 'normal', 0)

        if count_m < target_m:
            warnings.append({
                'type': 'staff_shortage', 'severity': 'high', 'day': day, 'shift': 'M',
                'message': f"روز {day}: کمبود پرسنل صبح ({count_m}/{target_m})",
            })
        if count_e < target_e:
            warnings.append({
                'type': 'staff_shortage', 'severity': 'high', 'day': day, 'shift': 'E',
                'message': f"روز {day}: کمبود پرسنل عصر ({count_e}/{target_e})",
            })
        if count_n < target_n:
            warnings.append({
                'type': 'staff_shortage', 'severity': 'high', 'day': day, 'shift': 'N',
                'message': f"روز {day}: کمبود پرسنل شب ({count_n}/{target_n})",
            })

        if count_m > target_m + 1:
            warnings.append({
                'type': 'staff_surplus', 'severity': 'low', 'day': day, 'shift': 'M',
                'message': f"روز {day}: اضافه پرسنل صبح ({count_m}/{target_m})",
            })
        if count_e > target_e + 1:
            warnings.append({
                'type': 'staff_surplus', 'severity': 'low', 'day': day, 'shift': 'E',
                'message': f"روز {day}: اضافه پرسنل عصر ({count_e}/{target_e})",
            })
        if count_n > target_n + 1:
            warnings.append({
                'type': 'staff_surplus', 'severity': 'low', 'day': day, 'shift': 'N',
                'message': f"روز {day}: اضافه پرسنل شب ({count_n}/{target_n})",
            })

    # ============================================================
    # 2. هشدارهای هر پرستار
    # ============================================================
    for profile in profiles:
        staff_id = profile.user.id
        staff_name = f"{profile.user.first_name} {profile.user.last_name}"
        staff_shifts = shifts_dict.get(staff_id, {})

        if not staff_shifts:
            warnings.append({
                'type': 'no_shifts', 'severity': 'high',
                'staff_id': staff_id, 'staff_name': staff_name,
                'message': f"{staff_name}: هیچ شیفتی در این ماه ندارد",
            })
            continue

        m_count = sum(1 for d in staff_shifts.values() if 'M' in d)
        e_count = sum(1 for d in staff_shifts.values() if 'E' in d)
        n_count = sum(1 for d in staff_shifts.values() if 'N' in d)
        off_count = sum(1 for d in staff_shifts.values() if d == '-')

        total_work = m_count + e_count + n_count

        # شب‌های متوالی
        nights = sorted([d for d, s in staff_shifts.items() if 'N' in s])
        for i in range(1, len(nights)):
            if nights[i] - nights[i-1] == 1:
                warnings.append({
                    'type': 'consecutive_nights', 'severity': 'high',
                    'staff_id': staff_id, 'staff_name': staff_name, 'day': nights[i],
                    'message': f"{staff_name}: شب متوالی در روزهای {nights[i-1]} و {nights[i]}",
                })

        # کار بعد از شب
        for day in range(2, days_in_month + 1):
            prev_shift = staff_shifts.get(day - 1, '')
            current_shift = staff_shifts.get(day, '')
            if 'N' in prev_shift and current_shift and current_shift != '-' and current_shift != 'H':
                warnings.append({
                    'type': 'work_after_night', 'severity': 'high',
                    'staff_id': staff_id, 'staff_name': staff_name, 'day': day,
                    'message': f"{staff_name}: کار بعد از شب — روز {day-1} شب، روز {day} شیفت {current_shift}",
                })

        # عدم تعادل شیفت
        if total_work > 5:
            avg = total_work / 3
            if m_count > avg * 2.5:
                warnings.append({
                    'type': 'shift_imbalance', 'severity': 'medium',
                    'staff_id': staff_id, 'staff_name': staff_name,
                    'message': f"{staff_name}: تعداد زیاد شیفت صبح ({m_count} از {total_work})",
                })
            if n_count > avg * 2.5:
                warnings.append({
                    'type': 'shift_imbalance', 'severity': 'medium',
                    'staff_id': staff_id, 'staff_name': staff_name,
                    'message': f"{staff_name}: تعداد زیاد شیفت شب ({n_count} از {total_work})",
                })
            if e_count > avg * 2.5:
                warnings.append({
                    'type': 'shift_imbalance', 'severity': 'medium',
                    'staff_id': staff_id, 'staff_name': staff_name,
                    'message': f"{staff_name}: تعداد زیاد شیفت عصر ({e_count} از {total_work})",
                })

        if off_count == 0 and total_work > 20:
            warnings.append({
                'type': 'no_off', 'severity': 'medium',
                'staff_id': staff_id, 'staff_name': staff_name,
                'message': f"{staff_name}: هیچ روز آفی در این ماه ندارد ({total_work} شیفت)",
            })

        if 0 < total_work < 10:
            warnings.append({
                'type': 'low_shifts', 'severity': 'medium',
                'staff_id': staff_id, 'staff_name': staff_name,
                'message': f"{staff_name}: شیفت‌های خیلی کم ({total_work} شیفت)",
            })

    return warnings


class MonthlyReportView(APIView):
    """
    گزارش تجمیعی چهار برنامه برای یک بخش در یک ماه.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        ward_id = request.query_params.get('ward_id')
        year = request.query_params.get('year')
        month = request.query_params.get('month')

        if not ward_id or not year or not month:
            return Response(
                {'error': 'ward_id، year و month الزامی هستند.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            ward = Ward.objects.get(id=ward_id)
        except Ward.DoesNotExist:
            return Response(
                {'error': 'بخش مورد نظر یافت نشد.'},
                status=status.HTTP_404_NOT_FOUND
            )

        # ✅ چک دسترسی
        user_profile = request.user.profile
        user_role = user_profile.role

        allowed_ward_ids = []
        if user_role == 'head':
            if user_profile.ward:
                allowed_ward_ids.append(user_profile.ward.id)
        elif user_role == 'supervisor':
            allowed_ward_ids = list(
                request.user.supervised_wards.values_list('id', flat=True)
            )
        elif user_role == 'admin':
            allowed_ward_ids = list(Ward.objects.values_list('id', flat=True))

        if ward.id not in allowed_ward_ids:
            return Response(
                {'error': 'شما به این بخش دسترسی ندارید.'},
                status=status.HTTP_403_FORBIDDEN
            )

        month_prefix = f"{year}-{str(month).zfill(2)}"
        days_in_month = get_persian_days_in_month(int(year), int(month))
        first_day_index = get_persian_weekday(int(year), int(month), 1)

        holidays = Holiday.objects.filter(
            date__startswith=month_prefix
        ).values('date', 'name')
        holidays_dict = {h['date']: h['name'] for h in holidays}

        result = {}

        # ✅ شامل assistant
        for st in ALL_SCHEDULE_TYPES:
            # ✅ استفاده از تابع کمکی برای پرسنل
            profiles = list(_get_profiles_for_schedule_type(ward, st, year, month))

            final_schedules = FinalSchedule.objects.filter(
                ward=ward,
                schedule_date__startswith=month_prefix,
                schedule_type=st
            ).select_related('assigned_staff', 'shift_code')

            total_m = 0
            total_e = 0
            total_n = 0
            total_h = 0
            total_off = 0
            staff_with_shifts = set()

            daily_m = [0] * days_in_month
            daily_e = [0] * days_in_month
            daily_n = [0] * days_in_month

            staff_stats = {}
            shifts_dict = {}

            for fs in final_schedules:
                staff_id = fs.assigned_staff.id
                day = int(fs.schedule_date.split('-')[2])
                shift_code = fs.shift_code.code if fs.shift_code else ''

                if not shift_code:
                    continue

                staff_with_shifts.add(staff_id)

                if staff_id not in shifts_dict:
                    shifts_dict[staff_id] = {}
                shifts_dict[staff_id][day] = shift_code

                if staff_id not in staff_stats:
                    staff_stats[staff_id] = {
                        'staff_id': staff_id,
                        'staff_name': f"{fs.assigned_staff.first_name} {fs.assigned_staff.last_name}",
                        'm': 0, 'e': 0, 'n': 0, 'h': 0, 'off': 0,
                        'total': 0
                    }

                if 'M' in shift_code:
                    total_m += 1
                    staff_stats[staff_id]['m'] += 1
                    daily_m[day - 1] += 1
                if 'E' in shift_code:
                    total_e += 1
                    staff_stats[staff_id]['e'] += 1
                    daily_e[day - 1] += 1
                if 'N' in shift_code:
                    total_n += 1
                    staff_stats[staff_id]['n'] += 1
                    daily_n[day - 1] += 1
                if 'H' in shift_code:
                    total_h += 1
                    staff_stats[staff_id]['h'] += 1
                if '-' in shift_code:
                    total_off += 1
                    staff_stats[staff_id]['off'] += 1

                staff_stats[staff_id]['total'] += 1

            # ✅ وضعیت
            schedule_status_obj = MonthlyScheduleStatus.objects.filter(
                ward=ward, year=int(year), month=int(month), schedule_type=st
            ).first()

            current_status = schedule_status_obj.status if schedule_status_obj else 'DRAFT'

            status_display = {
                'DRAFT': '⏳ پیش‌نویس',
                'PENDING_SUPERVISOR': '⏳ در انتظار تایید سوپروایزر',
                'PENDING_MATRON': '⏳ در انتظار تایید نهایی مترون',
                'LOCKED': '🔒 قفل شده',
                'RETURNED_TO_HEAD': '🔄 بازگشته به سرپرستار',
                'RETURNED_TO_SUPERVISOR': '🔄 بازگشته به سوپروایزر',
            }.get(current_status, 'نامشخص')

            # ✅ اهداف
            (
                target_m, target_e, target_n,
                holiday_target_m, holiday_target_e, holiday_target_n
            ) = _get_targets_for_schedule_type(ward, st)

            # ✅ پوشش اهداف
            total_required_m = 0
            total_required_e = 0
            total_required_n = 0

            for day in range(1, days_in_month + 1):
                day_index = first_day_index + (day - 1)
                date_key = f"{year}-{str(month).zfill(2)}-{str(day).zfill(2)}"
                is_holiday_or_friday = (date_key in holidays_dict) or (day_index % 7 == 6)

                total_required_m += holiday_target_m if is_holiday_or_friday else target_m
                total_required_e += holiday_target_e if is_holiday_or_friday else target_e
                total_required_n += holiday_target_n if is_holiday_or_friday else target_n

            coverage_m = int((total_m / total_required_m * 100)) if total_required_m > 0 else 100
            coverage_e = int((total_e / total_required_e * 100)) if total_required_e > 0 else 100
            coverage_n = int((total_n / total_required_n * 100)) if total_required_n > 0 else 100

            total_records = final_schedules.count()
            has_schedule = total_records > 0

            # ✅ محاسبه هشدارها
            daily_target_m = {'normal': target_m, 'holiday': holiday_target_m}
            daily_target_e = {'normal': target_e, 'holiday': holiday_target_e}
            daily_target_n = {'normal': target_n, 'holiday': holiday_target_n}

            warnings_list = []

            if has_schedule:
                warnings_list = analyze_schedule_warnings(
                    shifts_dict=shifts_dict,
                    profiles=profiles,
                    days_in_month=days_in_month,
                    first_day_index=first_day_index,
                    holidays_dict=holidays_dict,
                    year=int(year),
                    month=int(month),
                    ward=ward,
                    schedule_type=st,
                    daily_target_m=daily_target_m,
                    daily_target_e=daily_target_e,
                    daily_target_n=daily_target_n,
                )

            result[st] = {
                'exists': has_schedule,
                'status': current_status,
                'status_display': status_display,
                'staff_count': len(staff_with_shifts),
                'total_records': total_records,
                'total_m': total_m,
                'total_e': total_e,
                'total_n': total_n,
                'total_h': total_h,
                'total_off': total_off,
                'target_m': target_m,
                'target_e': target_e,
                'target_n': target_n,
                'holiday_target_m': holiday_target_m,
                'holiday_target_e': holiday_target_e,
                'holiday_target_n': holiday_target_n,
                'total_required_m': total_required_m,
                'total_required_e': total_required_e,
                'total_required_n': total_required_n,
                'coverage_m': coverage_m,
                'coverage_e': coverage_e,
                'coverage_n': coverage_n,
                'daily_m': daily_m,
                'daily_e': daily_e,
                'daily_n': daily_n,
                'staff_stats': list(staff_stats.values()),
                'warnings': warnings_list,
                'warnings_count': len(warnings_list),
                'warnings_high': len([w for w in warnings_list if w['severity'] == 'high']),
                'warnings_medium': len([w for w in warnings_list if w['severity'] == 'medium']),
            }

        active_programs = sum(1 for st in result.values() if st['exists'])
        locked_programs = sum(1 for st in result.values() if st['status'] == 'LOCKED')

        result['summary'] = {
            'ward_id': ward.id,
            'ward_name': ward.name,
            'year': int(year),
            'month': int(month),
            'month_name': ['', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور', 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند'][int(month)],
            'days_in_month': days_in_month,
            'first_day_index': first_day_index,
            'weekdays': ['شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه'],
            'active_programs': active_programs,
            'locked_programs': locked_programs,
            'total_staff': len(profiles),
        }

        return Response(result, status=status.HTTP_200_OK)


# ========================================================
# ویوهای بخش‌های سوپروایزر/مترون
# ========================================================
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_my_supervised_wards(request):
    """دریافت لیست بخش‌های تحت نظارت سوپروایزر."""
    try:
        profile = request.user.profile
    except Profile.DoesNotExist:
        return Response([])

    if profile.role != 'supervisor':
        return Response([])

    wards = request.user.supervised_wards.all()
    result = []
    for ward in wards:
        result.append({
            'id': ward.id,
            'name': ward.name,
        })

    return Response(result)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_all_wards_for_matron(request):
    """دریافت لیست همه بخش‌ها برای مترون."""
    try:
        profile = request.user.profile
    except Profile.DoesNotExist:
        return Response([])

    if profile.role != 'admin':
        return Response([])

    wards = Ward.objects.all().order_by('name')
    result = []
    for ward in wards:
        result.append({
            'id': ward.id,
            'name': ward.name,
        })

    return Response(result)


# ========================================================
# تغییر رمز و آپلود آواتار
# ========================================================
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError as DjangoValidationError


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def change_password(request):
    """تغییر رمز عبور توسط کاربر."""
    user = request.user
    old_password = request.data.get('old_password')
    new_password = request.data.get('new_password')
    confirm_password = request.data.get('confirm_password')

    if not old_password or not new_password or not confirm_password:
        return Response(
            {'error': 'همه فیلدها الزامی هستند.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    if not user.check_password(old_password):
        return Response(
            {'error': 'رمز عبور فعلی اشتباه است.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    if new_password != confirm_password:
        return Response(
            {'error': 'رمز جدید و تکرار آن یکسان نیستند.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    try:
        validate_password(new_password, user)
    except DjangoValidationError as e:
        return Response(
            {'error': ' '.join(e.messages)},
            status=status.HTTP_400_BAD_REQUEST
        )

    user.set_password(new_password)
    user.save()

    update_session_auth_hash(request, user)

    return Response({
        'status': 'success',
        'message': 'رمز عبور با موفقیت تغییر کرد.'
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def upload_avatar(request):
    """آپلود عکس پروفایل."""
    try:
        profile = request.user.profile
    except Profile.DoesNotExist:
        return Response(
            {'error': 'پروفایل کاربر یافت نشد.'},
            status=status.HTTP_404_NOT_FOUND
        )

    if 'avatar' not in request.FILES:
        return Response(
            {'error': 'فایل عکس ارسال نشده است.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    avatar = request.FILES['avatar']

    if avatar.size > 2 * 1024 * 1024:
        return Response(
            {'error': 'حجم عکس نباید بیشتر از ۲ مگابایت باشد.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    if not avatar.content_type.startswith('image/'):
        return Response(
            {'error': 'فقط فایل تصویری مجاز است.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    profile.avatar = avatar
    profile.save()

    return Response({
        'status': 'success',
        'message': 'عکس پروفایل با موفقیت آپلود شد.',
        'avatar_url': profile.avatar.url if profile.avatar else None,
    })


# ========================================================
# آپلود اکسل
# ========================================================
@api_view(['GET'])
@permission_classes([AllowAny])
def download_excel_template(request):
    """دانلود قالب نمونه اکسل — بدون نیاز به احراز هویت"""
    template_path = os.path.join(settings.BASE_DIR, 'نمونه_اطلاعات.xlsx')

    if not os.path.exists(template_path):
        try:
            import subprocess
            subprocess.run(
                ['python', 'create_sample_excel.py'],
                cwd=settings.BASE_DIR,
                check=True
            )
        except Exception as e:
            return Response(
                {'error': f'خطا در ساخت قالب: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    if not os.path.exists(template_path):
        return Response(
            {'error': 'قالب نمونه یافت نشد'},
            status=status.HTTP_404_NOT_FOUND
        )

    response = FileResponse(
        open(template_path, 'rb'),
        as_attachment=True,
        filename='نمونه_اطلاعات.xlsx'
    )
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


class ExcelUploadView(APIView):
    """آپلود فایل اکسل و پردازش خودکار"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        try:
            profile = request.user.profile
        except:
            return Response(
                {'error': 'پروفایل کاربر یافت نشد'},
                status=status.HTTP_404_NOT_FOUND
            )

        if profile.role not in ['admin', 'head']:
            return Response(
                {'error': 'شما اجازه آپلود فایل را ندارید'},
                status=status.HTTP_403_FORBIDDEN
            )

        if 'file' not in request.FILES:
            return Response(
                {'error': 'فایلی ارسال نشده است'},
                status=status.HTTP_400_BAD_REQUEST
            )

        uploaded_file = request.FILES['file']

        from .excel_import import validate_file
        is_valid, error_msg = validate_file(uploaded_file)

        if not is_valid:
            return Response(
                {'error': error_msg},
                status=status.HTTP_400_BAD_REQUEST
            )

        from .models import ExcelUpload
        from .excel_import import process_excel_upload

        excel_upload = ExcelUpload.objects.create(
            file=uploaded_file,
            uploaded_by=request.user,
            status='processing'
        )

        try:
            result = process_excel_upload(excel_upload)
        except Exception as e:
            import traceback
            traceback.print_exc()
            return Response(
                {'error': f'خطا در پردازش فایل: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

        summary = result.get('summary', {})

        sections = {}
        for key in ['staff', 'wards', 'holidays', 'shift_codes']:
            if result.get(key):
                sections[key] = {
                    'total': result[key].get('total', 0),
                    'created': result[key].get('created', 0),
                    'updated': result[key].get('updated', 0),
                    'errors': result[key].get('errors', []),
                }

        return Response({
            'status': 'success' if excel_upload.status == 'success' else 'partial' if excel_upload.status == 'partial' else 'failed',
            'upload_id': excel_upload.id,
            'is_success': excel_upload.status == 'success',
            'message': 'پردازش با موفقیت انجام شد' if excel_upload.status == 'success'
                       else 'پردازش با برخی خطاها انجام شد' if excel_upload.status == 'partial'
                       else 'پردازش ناموفق بود',
            'created_count': summary.get('total_created', 0),
            'updated_count': summary.get('total_updated', 0),
            'error_count': summary.get('total_errors', 0),
            'sections': sections,
        }, status=status.HTTP_200_OK)


class ExcelUploadHistoryView(APIView):
    """لیست تاریخچه آپلودها"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            profile = request.user.profile
        except:
            return Response([])

        if profile.role not in ['admin', 'head']:
            return Response([])

        from .models import ExcelUpload

        uploads = ExcelUpload.objects.all().order_by('-created_at')[:20]

        result = []
        for upload in uploads:
            summary = (upload.result_data or {}).get('summary', {})

            result.append({
                'id': upload.id,
                'file_name': os.path.basename(upload.file.name) if upload.file else 'نامشخص',
                'status': upload.status,
                'status_display': upload.get_status_display(),
                'created_count': summary.get('total_created', 0),
                'updated_count': summary.get('total_updated', 0),
                'error_count': summary.get('total_errors', 0),
                'uploaded_by_name': (
                    f"{upload.uploaded_by.first_name} {upload.uploaded_by.last_name}".strip()
                    if upload.uploaded_by else None
                ),
                'created_at': upload.created_at.isoformat(),
                'processed_at': upload.processed_at.isoformat() if upload.processed_at else None,
            })

        return Response(result)


# ========================================================
# مدیریت پرسنل برنامه
# ========================================================
from .models import ScheduleStaffOverride


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_ward_staff_with_overrides(request):
    """
    ✅ لیست پرسنل بخش با اعمال حذف‌های موقت (بر اساس schedule_type).
    """
    ward_id = request.query_params.get('ward_id')
    year = request.query_params.get('year')
    month = request.query_params.get('month')
    schedule_type = request.query_params.get('schedule_type', 'main')

    if not ward_id:
        return Response({'error': 'ward_id الزامی است'}, status=400)

    try:
        ward = Ward.objects.get(id=ward_id)
    except Ward.DoesNotExist:
        return Response({'error': 'بخش یافت نشد'}, status=404)

    # ✅ پرسنل پایه بر اساس schedule_type
    base_staff = list(_get_profiles_for_schedule_type(ward, schedule_type))
    base_ids = {p.user.id for p in base_staff}

    # حذف‌ها
    overrides = ScheduleStaffOverride.objects.filter(
        ward=ward,
        year=int(year) if year else 0,
        month=int(month) if month else 0,
        schedule_type=schedule_type,
    ).select_related('staff', 'created_by')

    excluded_ids = {o.staff_id for o in overrides}

    # پرسنل فعال
    active_staff = [p for p in base_staff if p.user.id not in excluded_ids]

    active_list = []
    for p in active_staff:
        active_list.append({
            'staff_id': p.user.id,
            'staff_name': f"{p.user.first_name} {p.user.last_name}".strip() or p.user.username,
            'position_label': p.get_position_display() if hasattr(p, 'get_position_display') else '',
            'employment_status': p.employment_status or '',
        })

    # پرسنل حذف‌شده
    removed_list = []
    for o in overrides:
        try:
            p = Profile.objects.get(user=o.staff)
            removed_list.append({
                'staff_id': o.staff.id,
                'staff_name': f"{o.staff.first_name} {o.staff.last_name}".strip() or o.staff.username,
                'position_label': p.get_position_display() if hasattr(p, 'get_position_display') else '',
                'reason': o.reason or '',
                'created_by_name': (
                    f"{o.created_by.first_name} {o.created_by.last_name}".strip()
                    if o.created_by else ''
                ),
                'created_at': o.created_at.isoformat() if o.created_at else None,
            })
        except Profile.DoesNotExist:
            continue

    return Response({
        'active_staff': active_list,
        'removed_staff': removed_list,
        'total_base': len(base_staff),
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def exclude_staff_from_schedule(request):
    """
    حذف موقت یک پرستار از برنامه.
    """
    user = request.user
    try:
        profile = user.profile
    except Profile.DoesNotExist:
        return Response({'error': 'پروفایل یافت نشد'}, status=404)

    if profile.role != 'head':
        return Response(
            {'error': 'فقط سرپرستار می‌تواند پرسنل را حذف کند.'},
            status=403
        )

    ward_id = request.data.get('ward_id')
    year = request.data.get('year')
    month = request.data.get('month')
    schedule_type = request.data.get('schedule_type', 'main')
    staff_id = request.data.get('staff_id')
    reason = request.data.get('reason', '')

    if not all([ward_id, year, month, staff_id]):
        return Response({'error': 'پارامترهای ناقص'}, status=400)

    try:
        ward = Ward.objects.get(id=ward_id)
        staff = User.objects.get(id=staff_id)
    except (Ward.DoesNotExist, User.DoesNotExist):
        return Response({'error': 'بخش یا پرستار یافت نشد'}, status=404)

    override, created = ScheduleStaffOverride.objects.update_or_create(
        ward=ward,
        year=int(year),
        month=int(month),
        schedule_type=schedule_type,
        staff=staff,
        defaults={
            'reason': reason,
            'created_by': user,
        }
    )

    return Response({
        'status': 'success',
        'created': created,
        'staff_name': f"{staff.first_name} {staff.last_name}".strip() or staff.username,
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def restore_staff_to_schedule(request):
    """
    بازگرداندن پرستار حذف‌شده به برنامه.
    """
    user = request.user
    try:
        profile = user.profile
    except Profile.DoesNotExist:
        return Response({'error': 'پروفایل یافت نشد'}, status=404)

    if profile.role != 'head':
        return Response(
            {'error': 'فقط سرپرستار می‌تواند پرسنل را بازگرداند.'},
            status=403
        )

    ward_id = request.data.get('ward_id')
    year = request.data.get('year')
    month = request.data.get('month')
    schedule_type = request.data.get('schedule_type', 'main')
    staff_id = request.data.get('staff_id')

    if not all([ward_id, year, month, staff_id]):
        return Response({'error': 'پارامترهای ناقص'}, status=400)

    deleted_count, _ = ScheduleStaffOverride.objects.filter(
        ward_id=ward_id,
        year=int(year),
        month=int(month),
        schedule_type=schedule_type,
        staff_id=staff_id,
    ).delete()

    if deleted_count == 0:
        return Response({'error': 'رکورد حذف یافت نشد'}, status=404)

    return Response({
        'status': 'restored',
        'count': deleted_count,
    })